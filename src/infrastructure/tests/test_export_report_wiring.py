from datetime import datetime
from os import environ
from os.path import isfile, join
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import RGBColor

from src.application.export_report_use_case import ExportReportUseCase
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures
from src.infrastructure.wirings.export_report_wiring import ExportReportWiring

SECTION_HEADINGS_IN_REPORT_ORDER = [
    "RESUMEN EJECUTIVO",
    "📄 INFORMACIÓN DEL DOCUMENTO",
    "🏷️ CLASIFICACIÓN DEL ARTÍCULO",
    "⭐ ANÁLISIS DE CALIDAD SEMÁNTICA",
    "🎯 PERTINENCIA EDITORIAL",
    "📝 GRAMÁTICA Y ORTOGRAFÍA",
    "📖 VALIDACIÓN APA 7 (ESPAÑOL)",
    "📋 VALIDACIÓN DE ESTRUCTURA (EUMIC)",
    "📚 ANÁLISIS DE CITAS Y REFERENCIAS",
    "💡 RECOMENDACIONES",
]


class TestExportReportWiring(TestCase):
    def setUp(self) -> None:
        self._temporary_directory = TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self._output_path = join(self._temporary_directory.name, "report.docx")

    def _export_with_environment(self, environment: dict[str, str]) -> DocxDocument:
        with patch.dict(environ, environment):
            use_case = ExportReportWiring().create_use_case()
            export_succeeded = use_case.execute(
                report_input=ReportFixtures.make_report_input_dto(),
                output_path=self._output_path,
            )
        self.assertTrue(export_succeeded)
        return Document(self._output_path)

    def _paragraph_texts(self, document: DocxDocument) -> list[str]:
        return [paragraph.text for paragraph in document.paragraphs]

    def _footer_stamp(self, document: DocxDocument) -> str:
        footer_stamps = [
            text for text in self._paragraph_texts(document) if text.startswith("Generado por")
        ]
        self.assertEqual(len(footer_stamps), 1)
        return footer_stamps[0]

    def test_create_use_case_returns_export_report_use_case_instance(self):
        result = ExportReportWiring().create_use_case()
        self.assertIsInstance(result, ExportReportUseCase)

    def test_execute_writes_docx_file_at_output_path(self):
        self._export_with_environment({})

        self.assertTrue(isfile(self._output_path))

    def test_execute_estimates_pages_with_report_words_per_page_from_env(self):
        document = self._export_with_environment({"REPORT_WORDS_PER_PAGE": "300"})

        self.assertIn("Páginas estimadas: 8", self._paragraph_texts(document))

    def test_execute_stamps_app_name_and_version_from_env(self):
        document = self._export_with_environment(
            {"TESTING": "True", "SILVINA_APP_NAME": "Custom App", "SILVINA_VERSION": "1.0"}
        )

        self.assertTrue(self._footer_stamp(document).startswith("Generado por Custom App v1.0 | "))

    def test_execute_colors_overall_score_with_score_thresholds_from_env(self):
        document = self._export_with_environment(
            {"REPORT_SCORE_HIGH_THRESHOLD": "9.0", "REPORT_SCORE_MEDIUM_THRESHOLD": "5.0"}
        )

        overall_score_paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("Puntuación General: ")
        )
        warning_color = ReportFixtures.make_settings().warning_color_rgb
        self.assertEqual(overall_score_paragraph.runs[1].font.color.rgb, RGBColor(*warning_color))

    def test_execute_decides_publishability_with_publish_threshold_from_env(self):
        document = self._export_with_environment({"PUBLISH_THRESHOLD": "8.0"})

        self.assertIn("⚠️ REQUIERE REVISIÓN", self._paragraph_texts(document))

    def test_execute_stamps_footer_with_current_date_from_system_clock(self):
        document = self._export_with_environment({})

        self.assertIn(datetime.now().strftime("%d/%m/%Y"), self._footer_stamp(document))

    def test_execute_decorates_header_with_logo_and_current_date_from_system_clock(self):
        document = self._export_with_environment({})

        header = document.sections[0].header
        header_relationship_types = [
            relationship.reltype for relationship in header.part.rels.values()
        ]
        self.assertEqual(header_relationship_types.count(RELATIONSHIP_TYPE.IMAGE), 1)
        self.assertTrue(
            header.tables[0]
            .rows[0]
            .cells[0]
            .text.endswith(datetime.now().strftime("%d de %B de %Y"))
        )

    def test_execute_decorates_footer_with_page_and_total_pages_fields(self):
        document = self._export_with_environment({})

        footer_paragraph = document.sections[0].footer.paragraphs[0]
        field_instructions = footer_paragraph._element.xpath(".//w:instrText")
        self.assertEqual(
            [instruction.text for instruction in field_instructions], ["PAGE", "NUMPAGES"]
        )

    def test_execute_renders_every_section_heading_in_report_order(self):
        document = self._export_with_environment({})

        section_headings = []
        for paragraph in document.paragraphs:
            paragraph_style = paragraph.style
            assert paragraph_style is not None
            if paragraph_style.name == "Heading 1":
                section_headings.append(paragraph.text)
        self.assertEqual(section_headings, SECTION_HEADINGS_IN_REPORT_ORDER)

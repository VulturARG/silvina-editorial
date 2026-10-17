from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from src.domain.exceptions.document_errors import DocumentNotFound, DocumentUnreadable
from src.infrastructure.adapters.document.docx_table_adapter import DocxTableAdapter


class TestDocxTableAdapter(TestCase):
    def setUp(self):
        self.adapter = DocxTableAdapter()
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)

    def test_table_position_counts_only_non_empty_paragraphs_before_it(self):
        document = Document()
        document.add_paragraph("Introduction")
        document.add_paragraph("Methods")
        self._add_table(document, [["Name", "Value"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].position, 2)
        self.assertEqual(tables[0].rows, (("Name", "Value"),))

    def test_empty_paragraphs_before_table_are_not_counted(self):
        document = Document()
        document.add_paragraph("Introduction")
        document.add_paragraph("")
        document.add_paragraph("   ")
        self._add_table(document, [["Name"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].position, 1)

    def test_each_table_gets_its_own_position(self):
        document = Document()
        document.add_paragraph("Introduction")
        self._add_table(document, [["first"]])
        document.add_paragraph("Discussion")
        document.add_paragraph("Conclusion")
        self._add_table(document, [["second"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual([table.position for table in tables], [1, 3])
        self.assertEqual([table.rows for table in tables], [(("first",),), (("second",),)])

    def test_horizontally_merged_cell_is_not_repeated(self):
        document = Document()
        table = self._add_table(document, [["Merged", ""], ["a", "b"]])
        table.cell(0, 0).merge(table.cell(0, 1))

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].rows, (("Merged",), ("a", "b")))

    def test_vertically_merged_cell_is_not_repeated_in_later_rows(self):
        document = Document()
        table = self._add_table(document, [["Group", "x"], ["", "y"], ["", "z"]])
        table.cell(0, 0).merge(table.cell(2, 0))

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].rows, (("Group", "x"), ("", "y"), ("", "z")))

    def test_nested_table_text_is_ignored(self):
        document = Document()
        table = self._add_table(document, [["Outer", "b"]])
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Inner"

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].rows, (("Outer", "b"),))

    def test_cell_text_is_stripped(self):
        document = Document()
        self._add_table(document, [["  padded  ", "value"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].rows, (("padded", "value"),))

    def test_all_empty_table_is_skipped(self):
        document = Document()
        document.add_paragraph("Introduction")
        self._add_table(document, [["", ""], ["", ""]])
        document.add_paragraph("Discussion")
        self._add_table(document, [["kept"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0].position, 2)
        self.assertEqual(tables[0].rows, (("kept",),))

    def test_rows_with_only_empty_cells_are_dropped(self):
        document = Document()
        self._add_table(document, [["Name", "Value"], ["", "   "], ["Alpha", "1"]])

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables[0].rows, (("Name", "Value"), ("Alpha", "1")))

    def test_document_without_tables_returns_empty_list(self):
        document = Document()
        document.add_paragraph("Only text")

        tables = self.adapter.read_tables(path=self._save(document))

        self.assertEqual(tables, [])

    def test_missing_file_raises_document_not_found(self):
        missing_path = Path(self.temporary_directory.name) / "missing.docx"

        with self.assertRaises(DocumentNotFound):
            self.adapter.read_tables(path=str(missing_path))

    def test_corrupt_file_raises_document_unreadable(self):
        corrupt_path = Path(self.temporary_directory.name) / "corrupt.docx"
        corrupt_path.write_bytes(b"not a valid docx package")

        with self.assertRaises(DocumentUnreadable):
            self.adapter.read_tables(path=str(corrupt_path))

    def _add_table(self, document, rows):
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for column_index, text in enumerate(row):
                table.cell(row_index, column_index).text = text
        return table

    def _save(self, document, file_name="document.docx"):
        document_path = Path(self.temporary_directory.name) / file_name
        document.save(str(document_path))
        return str(document_path)

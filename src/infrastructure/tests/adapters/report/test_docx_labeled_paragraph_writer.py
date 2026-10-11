from unittest import TestCase

from docx import Document

from src.infrastructure.adapters.report.docx_labeled_paragraph_writer import (
    DocxLabeledParagraphWriter,
)


class TestDocxLabeledParagraphWriter(TestCase):
    def setUp(self) -> None:
        self.writer = DocxLabeledParagraphWriter()
        self.document = Document()

    def test_write_creates_paragraph_with_bold_label_and_plain_value(self) -> None:
        paragraph = self.writer.write(
            document=self.document,
            label="Título: ",
            value="Mi Artículo de Prueba",
        )

        self.assertEqual(paragraph.text, "Título: Mi Artículo de Prueba")
        self.assertEqual(len(self.document.paragraphs), 1)
        self.assertEqual(len(paragraph.runs), 2)

        label_run = paragraph.runs[0]
        self.assertEqual(label_run.text, "Título: ")
        self.assertTrue(label_run.bold)

        value_run = paragraph.runs[1]
        self.assertEqual(value_run.text, "Mi Artículo de Prueba")
        self.assertFalse(value_run.bold)

    def test_write_preserves_label_text_unchanged(self) -> None:
        paragraph = self.writer.write(
            document=self.document,
            label="Total de palabras: ",
            value="12,500",
        )

        self.assertEqual(paragraph.runs[0].text, "Total de palabras: ")
        self.assertEqual(paragraph.runs[1].text, "12,500")

    def test_write_returns_created_paragraph(self) -> None:
        paragraph = self.writer.write(
            document=self.document,
            label="Estado: ",
            value="Correcto",
        )

        self.assertEqual(paragraph.text, self.document.paragraphs[0].text)
        self.assertIs(paragraph._element, self.document.paragraphs[0]._element)

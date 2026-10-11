from unittest import TestCase

from docx import Document

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxDocumentInfoSection(TestCase):
    def setUp(self) -> None:
        self.section = ReportFixtures.make_document_info_section()
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_section_heading(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(self.document.paragraphs[0].text, "📄 INFORMACIÓN DEL DOCUMENTO")

    def test_render_adds_labeled_document_metadata(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(
            self._paragraph_texts()[1:],
            [
                "Título: Test",
                "Autor(es): Test Author",
                "Total de palabras: 2,500",
                "Total de caracteres: 12,000",
                "Páginas estimadas: 10",
            ],
        )

    def test_render_falls_back_to_filename_and_unknown_authors(self) -> None:
        document_content = ReportFixtures.make_document_content_mock(title="")
        document_content.authors = ""
        report_input = ReportFixtures.make_report_input_dto(
            filename="manuscrito.docx", document_content=document_content
        )

        self.section.render(document=self.document, report_input=report_input)

        self.assertIn("Título: manuscrito.docx", self._paragraph_texts())
        self.assertIn("Autor(es): No identificado", self._paragraph_texts())

    def test_render_derives_estimated_pages_from_word_count(self) -> None:
        document_content = ReportFixtures.make_document_content_mock(word_count=500)
        report_input = ReportFixtures.make_report_input_dto(document_content=document_content)

        self.section.render(document=self.document, report_input=report_input)

        self.assertIn("Páginas estimadas: 2", self._paragraph_texts())

    def test_estimated_pages_uses_settings_words_per_page(self) -> None:
        settings = ReportFixtures.make_settings(words_per_page=100)
        section = ReportFixtures.make_document_info_section(docx_report_settings=settings)
        document_content = ReportFixtures.make_document_content_mock(word_count=250)
        report_input = ReportFixtures.make_report_input_dto(document_content=document_content)

        section.render(document=self.document, report_input=report_input)

        pages_text = [text for text in self._paragraph_texts() if "Páginas estimadas" in text]
        self.assertEqual(pages_text[0], "Páginas estimadas: 2")

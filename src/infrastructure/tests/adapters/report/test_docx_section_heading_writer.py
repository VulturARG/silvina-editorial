from unittest import TestCase

from docx import Document
from docx.shared import RGBColor

from src.infrastructure.adapters.report.docx_section_heading_writer import DocxSectionHeadingWriter
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxSectionHeadingWriter(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.writer = DocxSectionHeadingWriter(docx_report_settings=self.settings)
        self.document = Document()

    def test_write_creates_heading_with_default_level_one(self) -> None:
        heading = self.writer.write(
            document=self.document,
            text="RESUMEN EJECUTIVO",
        )

        self.assertEqual(heading.text, "RESUMEN EJECUTIVO")
        self.assertEqual(len(self.document.paragraphs), 1)
        self.assertIsNotNone(heading.style)
        assert heading.style is not None
        self.assertEqual(heading.style.name, "Heading 1")

    def test_write_applies_heading_color_to_all_runs(self) -> None:
        heading = self.writer.write(
            document=self.document,
            text="INFORMACIÓN DEL DOCUMENTO",
        )

        expected_color = RGBColor(*self.settings.heading_color_rgb)
        self.assertGreater(len(heading.runs), 0)
        for run in heading.runs:
            self.assertEqual(run.font.color.rgb, expected_color)

    def test_write_accepts_custom_heading_level(self) -> None:
        heading = self.writer.write(
            document=self.document,
            text="Subsección",
            level=2,
        )

        self.assertIsNotNone(heading.style)
        assert heading.style is not None
        self.assertEqual(heading.style.name, "Heading 2")
        expected_color = RGBColor(*self.settings.heading_color_rgb)
        for run in heading.runs:
            self.assertEqual(run.font.color.rgb, expected_color)

    def test_write_returns_heading_paragraph(self) -> None:
        heading = self.writer.write(
            document=self.document,
            text="Encabezado retornado",
        )

        self.assertEqual(heading.text, self.document.paragraphs[0].text)
        self.assertIs(heading._element, self.document.paragraphs[0]._element)

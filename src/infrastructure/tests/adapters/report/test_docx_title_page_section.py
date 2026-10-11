from unittest import TestCase

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxTitlePageSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_title_page_section(docx_report_settings=self.settings)
        self.document = Document()

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_centered_report_title_with_title_styling(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        title = self.document.paragraphs[0]
        self.assertEqual(title.text, "INFORME DE ANÁLISIS EDITORIAL")
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        for run in title.runs:
            self.assertEqual(run.font.size, Pt(self.settings.title_font_size_pt))
            self.assertEqual(run.font.color.rgb, RGBColor(*self.settings.heading_color_rgb))
            self.assertTrue(run.font.bold)
            self.assertTrue(run.font.underline)

    def test_render_adds_blank_paragraph_after_report_title(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(self.document.paragraphs[1].text, "")

    def test_render_adds_bold_centered_document_title(self) -> None:
        document_content = ReportFixtures.make_document_content_mock(title="Mi artículo")
        report_input = ReportFixtures.make_report_input_dto(document_content=document_content)

        self.section.render(document=self.document, report_input=report_input)

        document_title = self.document.paragraphs[2]
        self.assertEqual(document_title.text, "Mi artículo")
        self.assertEqual(document_title.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(document_title.runs[0].bold)
        self.assertEqual(document_title.runs[0].font.size, Pt(self.settings.base_font_size_pt))

    def test_render_falls_back_to_filename_when_document_title_is_empty(self) -> None:
        document_content = ReportFixtures.make_document_content_mock(title="")
        report_input = ReportFixtures.make_report_input_dto(
            filename="manuscrito.docx", document_content=document_content
        )

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(self.document.paragraphs[2].text, "manuscrito.docx")

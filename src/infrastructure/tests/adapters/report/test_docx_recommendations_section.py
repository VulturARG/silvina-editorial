from unittest import TestCase

from docx import Document
from docx.shared import Pt, RGBColor

from src.domain.dtos.recommendation_dto import RecommendationDTO
from src.domain.enums.publication_verdict import PublicationVerdict
from src.domain.enums.recommendation_priority import RecommendationPriority
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxRecommendationsSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_recommendations_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_bold_verdict_message_without_recommendations(self) -> None:
        report_input = ReportFixtures.make_report_input_dto()

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._paragraph_texts(), ["💡 RECOMENDACIONES", report_input.verdict.message]
        )
        verdict_run = self.document.paragraphs[1].runs[0]
        self.assertIs(verdict_run.bold, True)
        self.assertEqual(verdict_run.font.size, Pt(self.settings.recommendation_font_size_pt))

    def test_render_colors_verdict_message_by_verdict(self) -> None:
        cases = (
            (PublicationVerdict.CRITICAL, self.settings.reject_color_rgb),
            (PublicationVerdict.WARNING, self.settings.reject_color_rgb),
            (PublicationVerdict.APPROVED, self.settings.publishable_color_rgb),
        )
        for verdict, expected_color in cases:
            with self.subTest(verdict=verdict):
                document = Document()
                report_input = ReportFixtures.make_report_input_dto(
                    verdict=ReportFixtures.make_verdict_dto(verdict=verdict)
                )

                self.section.render(document=document, report_input=report_input)

                verdict_run = document.paragraphs[1].runs[0]
                self.assertEqual(verdict_run.font.color.rgb, RGBColor(*expected_color))

    def test_render_lists_recommendations_with_priority_icons_under_bold_label(self) -> None:
        recommendations = [
            RecommendationDTO(priority=RecommendationPriority.HIGH, message="Alta"),
            RecommendationDTO(priority=RecommendationPriority.MEDIUM, message="Media"),
            RecommendationDTO(priority=RecommendationPriority.LOW, message="Baja"),
        ]
        report_input = ReportFixtures.make_report_input_dto(recommendations=recommendations)

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._paragraph_texts()[2:],
            ["Recomendaciones específicas:", "🔴 Alta", "🟡 Media", "🟢 Baja"],
        )
        self.assertIs(self.document.paragraphs[2].runs[0].bold, True)
        bullet_style = self.document.paragraphs[3].style
        assert bullet_style is not None
        self.assertEqual(bullet_style.name, "List Bullet")

from re import compile as regex_compile
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

_HEADER_DATE_PATTERN = regex_compile(r"\b\d{1,2}\s+de\s+[^\W\d_]+\s+de\s+\d{4}\b")
_FOOTER_DATE_PATTERN = regex_compile(r"\b\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2}\b")
_ENGLISH_METRIC_UNITS_PER_INCH = 914400


class DocxFingerprintBuilder:
    """Builds a deterministic, JSON-serializable fingerprint of a Word document."""

    def build(self, docx_path: str) -> dict[str, Any]:
        """Read a docx file and return a structured fingerprint dictionary."""
        document = Document(docx_path)
        return {
            "body": self._extract_body(document),
            "footer": self._extract_footer(document),
            "header": self._extract_header(document),
            "normal_style": self._extract_normal_style(document),
        }

    def _extract_normal_style(self, document: DocxDocument) -> dict[str, Any]:
        style = document.styles["Normal"]
        font_size_points = (
            float(round(style.font.size.pt, 2)) if style.font.size is not None else None
        )
        space_before_points = (
            float(round(style.paragraph_format.space_before.pt, 2))
            if style.paragraph_format.space_before is not None
            else None
        )
        space_after_points = (
            float(round(style.paragraph_format.space_after.pt, 2))
            if style.paragraph_format.space_after is not None
            else None
        )
        return {
            "font_name": style.font.name,
            "font_size_pt": font_size_points,
            "line_spacing": style.paragraph_format.line_spacing,
            "space_after_pt": space_after_points,
            "space_before_pt": space_before_points,
        }

    def _extract_header(self, document: DocxDocument) -> dict[str, Any]:
        section = document.sections[0]
        header = section.header

        header_paragraphs = []
        for paragraph in header.paragraphs:
            header_paragraphs.append(self._extract_header_paragraph(paragraph))

        header_tables = []
        for table in header.tables:
            table_rows = []
            for row in table.rows:
                row_cells = []
                for cell in row.cells:
                    cell_paragraphs = [
                        self._extract_header_paragraph(cell_paragraph)
                        for cell_paragraph in cell.paragraphs
                    ]
                    row_cells.append(
                        {
                            "paragraphs": cell_paragraphs,
                            "text": self._mask_dates(cell.text),
                        }
                    )
                table_rows.append(row_cells)
            header_tables.append(
                {
                    "rows": table_rows,
                    "style_name": table.style.name if table.style is not None else None,
                }
            )

        return {
            "paragraphs": header_paragraphs,
            "tables": header_tables,
        }

    def _extract_header_paragraph(self, paragraph: Paragraph) -> dict[str, Any]:
        inline_elements = paragraph._element.xpath(".//wp:inline")
        has_inline_picture = len(inline_elements) > 0
        picture_width_inches = None
        if has_inline_picture:
            extent_elements = inline_elements[0].xpath(".//wp:extent")
            if extent_elements:
                extent_cx = extent_elements[0].get("cx")
                if extent_cx is not None:
                    picture_width_inches = float(
                        round(int(extent_cx) / _ENGLISH_METRIC_UNITS_PER_INCH, 4)
                    )

        alignment_name = paragraph.alignment.name if paragraph.alignment is not None else None
        return {
            "alignment": alignment_name,
            "has_inline_picture": has_inline_picture,
            "picture_width_inches": picture_width_inches,
            "text": self._mask_dates(paragraph.text),
        }

    def _extract_footer(self, document: DocxDocument) -> dict[str, Any]:
        section = document.sections[0]
        footer = section.footer

        footer_paragraphs = []
        for paragraph in footer.paragraphs:
            field_instructions = [
                instruction_node.text.strip()
                for instruction_node in paragraph._element.xpath(".//w:instrText")
            ]
            alignment_name = paragraph.alignment.name if paragraph.alignment is not None else None
            footer_paragraphs.append(
                {
                    "alignment": alignment_name,
                    "field_instructions": field_instructions,
                    "text": self._mask_dates(paragraph.text),
                }
            )

        return {
            "paragraphs": footer_paragraphs,
        }

    def _extract_body(self, document: DocxDocument) -> list[dict[str, Any]]:
        body_elements = []
        for child_element in document.element.body:
            element_tag = child_element.tag.split("}")[-1]
            if element_tag == "p":
                paragraph = Paragraph(child_element, document)
                body_elements.append(self._extract_body_paragraph(paragraph))
            elif element_tag == "tbl":
                table = Table(child_element, document)
                body_elements.append(self._extract_body_table(table))
        return body_elements

    def _extract_body_paragraph(self, paragraph: Paragraph) -> dict[str, Any]:
        left_indent = paragraph.paragraph_format.left_indent
        left_indent_points = float(round(left_indent.pt, 2)) if left_indent is not None else None
        alignment_name = paragraph.alignment.name if paragraph.alignment is not None else None
        style_name = paragraph.style.name if paragraph.style is not None else None
        runs_data = []
        for run in paragraph.runs:
            font_size_points = (
                float(round(run.font.size.pt, 2)) if run.font.size is not None else None
            )
            font_color_rgb = str(run.font.color.rgb) if run.font.color.rgb is not None else None
            runs_data.append(
                {
                    "bold": run.bold,
                    "font_color_rgb": font_color_rgb,
                    "font_size_pt": font_size_points,
                    "italic": run.italic,
                    "text": self._mask_dates(run.text),
                    "underline": run.underline,
                }
            )
        return {
            "alignment": alignment_name,
            "left_indent_pt": left_indent_points,
            "runs": runs_data,
            "style_name": style_name,
            "text": self._mask_dates(paragraph.text),
            "type": "paragraph",
        }

    def _extract_body_table(self, table: Table) -> dict[str, Any]:
        style_name = table.style.name if table.style is not None else None
        table_rows = []
        for row in table.rows:
            row_cells = [self._mask_dates(cell.text) for cell in row.cells]
            table_rows.append(row_cells)
        return {
            "rows": table_rows,
            "style_name": style_name,
            "type": "table",
        }

    def _mask_dates(self, text: str) -> str:
        if not text:
            return text
        header_date_masked = _HEADER_DATE_PATTERN.sub("<DATE>", text)
        return _FOOTER_DATE_PATTERN.sub("<DATE>", header_date_masked)

from unittest import TestCase

from docx import Document
from docx.shared import RGBColor

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxStructureValidationSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_structure_validation_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_reports_valid_structure_in_publishable_color(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(
            self._paragraph_texts(),
            [
                "📋 VALIDACIÓN DE ESTRUCTURA (EUMIC)",
                "✓ Estructura válida según normas EUMIC",
            ],
        )
        self.assertEqual(
            self.document.paragraphs[1].runs[0].font.color.rgb,
            RGBColor(*self.settings.publishable_color_rgb),
        )

    def test_render_reports_incomplete_structure_in_reject_color(self) -> None:
        structure = ReportFixtures.make_structure_mock(is_valid=False)
        report_input = ReportFixtures.make_report_input_dto(structure=structure)

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._paragraph_texts(),
            ["📋 VALIDACIÓN DE ESTRUCTURA (EUMIC)", "✗ Estructura incompleta"],
        )
        self.assertEqual(
            self.document.paragraphs[1].runs[0].font.color.rgb,
            RGBColor(*self.settings.reject_color_rgb),
        )

    def test_render_lists_missing_sections_when_structure_is_incomplete(self) -> None:
        structure = ReportFixtures.make_structure_mock(is_valid=False)
        structure.missing_sections = ["Resumen", "Conclusiones"]
        report_input = ReportFixtures.make_report_input_dto(structure=structure)

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._paragraph_texts()[2:],
            ["", "Secciones faltantes:", "• Resumen", "• Conclusiones"],
        )
        self.assertIs(self.document.paragraphs[3].runs[0].bold, True)
        bullet_style = self.document.paragraphs[4].style
        assert bullet_style is not None
        self.assertEqual(bullet_style.name, "List Bullet")

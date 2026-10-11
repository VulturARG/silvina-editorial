from unittest import TestCase

from docx import Document

from src.domain.enums.classification_confidence import ClassificationConfidence
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxClassificationSection(TestCase):
    def setUp(self) -> None:
        self.section = ReportFixtures.make_classification_section()
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_section_heading(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(self.document.paragraphs[0].text, "🏷️ CLASIFICACIÓN DEL ARTÍCULO")

    def test_render_reads_article_type_attribute_in_upper_case(self) -> None:
        classification = ReportFixtures.make_classification_mock(article_type_value="investigacion")
        report_input = ReportFixtures.make_report_input_dto(classification=classification)

        self.section.render(document=self.document, report_input=report_input)

        self.assertIn("Categoría: INVESTIGACION", self._paragraph_texts())

    def test_render_formats_classification_confidence_as_percentage(self) -> None:
        classification = ReportFixtures.make_classification_mock()
        classification.confidence = ClassificationConfidence.FULL_SIGNAL_MATCH
        report_input = ReportFixtures.make_report_input_dto(classification=classification)

        self.section.render(document=self.document, report_input=report_input)

        self.assertIn("Confianza: 90.0%", self._paragraph_texts())

    def test_render_shows_placeholder_when_confidence_is_missing(self) -> None:
        classification = ReportFixtures.make_classification_mock()
        classification.confidence = None
        report_input = ReportFixtures.make_report_input_dto(classification=classification)

        self.section.render(document=self.document, report_input=report_input)

        self.assertIn("Confianza: —", self._paragraph_texts())

    def test_render_adds_reasoning_when_present(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(self._paragraph_texts()[-1], "Razonamiento: Test reasoning")

    def test_render_omits_reasoning_when_empty(self) -> None:
        classification = ReportFixtures.make_classification_mock()
        classification.reasoning = ""
        report_input = ReportFixtures.make_report_input_dto(classification=classification)

        self.section.render(document=self.document, report_input=report_input)

        self.assertFalse(any(text.startswith("Razonamiento") for text in self._paragraph_texts()))

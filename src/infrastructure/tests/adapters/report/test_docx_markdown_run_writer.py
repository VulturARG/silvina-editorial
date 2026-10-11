from unittest import TestCase

from docx import Document

from src.infrastructure.adapters.report.docx_markdown_run_writer import DocxMarkdownRunWriter


class TestDocxMarkdownRunWriter(TestCase):
    def setUp(self) -> None:
        self.writer = DocxMarkdownRunWriter()
        self.document = Document()

    def test_write_paragraph_adds_plain_text_to_document(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Texto plano sin formato",
        )

        self.assertEqual(paragraph.text, "Texto plano sin formato")
        self.assertEqual(len(self.document.paragraphs), 1)
        self.assertFalse(any(run.bold for run in paragraph.runs))
        self.assertFalse(any(run.italic for run in paragraph.runs))

    def test_write_paragraph_renders_bold_runs_without_asterisks(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Inicio **texto en negrita** final",
        )

        self.assertEqual(paragraph.text, "Inicio texto en negrita final")
        self.assertNotIn("**", paragraph.text)
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        self.assertEqual(bold_runs, ["texto en negrita"])

    def test_write_paragraph_renders_italic_runs_without_asterisks(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="El texto emplea *correctamente* el *mecanismo* del *cual* depende.",
        )

        self.assertNotIn("*", paragraph.text)
        italic_runs = [run.text for run in paragraph.runs if run.italic]
        self.assertEqual(italic_runs, ["correctamente", "mecanismo", "cual"])

    def test_write_paragraph_renders_mixed_bold_and_italic_segments(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Texto con **negrita** y *cursiva*, más **negrita con *cursiva anidada***.",
        )

        self.assertNotIn("*", paragraph.text)
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        self.assertIn("negrita", bold_runs)
        self.assertIn("cursiva anidada", bold_runs)

        italic_runs = [run.text for run in paragraph.runs if run.italic]
        self.assertIn("cursiva", italic_runs)
        self.assertIn("cursiva anidada", italic_runs)

    def test_write_paragraph_preserves_lone_asterisks(self) -> None:
        text = "Operación a * b, lista '* item', escala 1 x 10^11, nota* y *unpaired."
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text=text,
        )

        self.assertIn("a * b", paragraph.text)
        self.assertIn("'* item'", paragraph.text)
        self.assertIn("1 x 10^11", paragraph.text)
        self.assertIn("nota*", paragraph.text)
        self.assertIn("*unpaired", paragraph.text)

    def test_write_paragraph_applies_prefix_to_target_paragraph(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Contenido del elemento",
            prefix="• ",
        )

        self.assertTrue(paragraph.text.startswith("• "))
        self.assertEqual(paragraph.runs[0].text, "• ")
        self.assertFalse(paragraph.runs[0].bold)

    def test_write_paragraph_uses_provided_existing_paragraph(self) -> None:
        existing_paragraph = self.document.add_paragraph()
        existing_paragraph.add_run("Prefijo existente: ")

        returned_paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Texto añadido con **negrita**",
            paragraph=existing_paragraph,
        )

        self.assertIs(returned_paragraph, existing_paragraph)
        self.assertEqual(len(self.document.paragraphs), 1)
        self.assertTrue(returned_paragraph.text.startswith("Prefijo existente: "))

    def test_write_paragraph_returns_created_paragraph(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Retorno verificado",
        )

        self.assertEqual(paragraph.text, "Retorno verificado")
        self.assertIs(paragraph._element, self.document.paragraphs[0]._element)

    def test_write_paragraph_handles_empty_text_without_creating_runs(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="",
        )

        self.assertEqual(paragraph.text, "")
        self.assertEqual(len(paragraph.runs), 0)

    def test_write_paragraph_with_empty_prefix_does_not_add_empty_run(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="Texto directo",
            prefix="",
        )

        self.assertEqual(len(paragraph.runs), 1)
        self.assertEqual(paragraph.runs[0].text, "Texto directo")

    def test_write_paragraph_renders_bold_only_without_italic(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="**solo negrita**",
        )

        self.assertEqual(len(paragraph.runs), 1)
        self.assertEqual(paragraph.runs[0].text, "solo negrita")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertIsNone(paragraph.runs[0].italic)

    def test_write_paragraph_renders_italic_only_without_bold(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="*solo cursiva*",
        )

        self.assertEqual(len(paragraph.runs), 1)
        self.assertEqual(paragraph.runs[0].text, "solo cursiva")
        self.assertTrue(paragraph.runs[0].italic)
        self.assertIsNone(paragraph.runs[0].bold)

    def test_write_paragraph_handles_bold_containing_italic(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="**prefijo *anidado* sufijo**",
        )

        self.assertEqual(len(paragraph.runs), 3)
        self.assertEqual(paragraph.runs[0].text, "prefijo ")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertIsNone(paragraph.runs[0].italic)

        self.assertEqual(paragraph.runs[1].text, "anidado")
        self.assertTrue(paragraph.runs[1].bold)
        self.assertTrue(paragraph.runs[1].italic)

        self.assertEqual(paragraph.runs[2].text, " sufijo")
        self.assertTrue(paragraph.runs[2].bold)
        self.assertIsNone(paragraph.runs[2].italic)

    def test_write_paragraph_handles_adjacent_bold_matches(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="**primer segmento** **segundo segmento**",
        )

        self.assertEqual(len(paragraph.runs), 3)
        self.assertEqual(paragraph.runs[0].text, "primer segmento")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertEqual(paragraph.runs[1].text, " ")
        self.assertIsNone(paragraph.runs[1].bold)
        self.assertEqual(paragraph.runs[2].text, "segundo segmento")
        self.assertTrue(paragraph.runs[2].bold)

    def test_write_paragraph_handles_adjacent_italic_matches(self) -> None:
        paragraph = self.writer.write_paragraph(
            document=self.document,
            text="*primer segmento* *segundo segmento*",
        )

        self.assertEqual(len(paragraph.runs), 3)
        self.assertEqual(paragraph.runs[0].text, "primer segmento")
        self.assertTrue(paragraph.runs[0].italic)
        self.assertEqual(paragraph.runs[1].text, " ")
        self.assertIsNone(paragraph.runs[1].italic)
        self.assertEqual(paragraph.runs[2].text, "segundo segmento")
        self.assertTrue(paragraph.runs[2].italic)

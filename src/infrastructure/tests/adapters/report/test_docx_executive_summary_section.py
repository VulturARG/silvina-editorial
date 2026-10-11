from unittest import TestCase

from docx import Document
from docx.shared import Pt, RGBColor

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxExecutiveSummarySection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_executive_summary_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _table_rows(self) -> list[list[str]]:
        return [[cell.text for cell in row.cells] for row in self.document.tables[0].rows]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_reports_publishable_decision_in_publishable_color(self) -> None:
        report_input = ReportFixtures.make_report_input_dto()

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            [paragraph.text for paragraph in self.document.paragraphs[:3]],
            [
                "RESUMEN EJECUTIVO",
                "✅ APTO PARA PUBLICACIÓN",
                report_input.publishability_reason,
            ],
        )
        decision_run = self.document.paragraphs[1].runs[0]
        self.assertIs(decision_run.bold, True)
        self.assertEqual(decision_run.font.size, Pt(self.settings.decision_font_size_pt))
        self.assertEqual(
            decision_run.font.color.rgb, RGBColor(*self.settings.publishable_color_rgb)
        )

    def test_render_reports_review_decision_in_reject_color(self) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            structure=ReportFixtures.make_structure_mock(is_valid=False)
        )

        self.section.render(document=self.document, report_input=report_input)

        decision_run = self.document.paragraphs[1].runs[0]
        self.assertEqual(decision_run.text, "⚠️ REQUIERE REVISIÓN")
        self.assertEqual(decision_run.font.color.rgb, RGBColor(*self.settings.reject_color_rgb))

    def test_render_adds_metrics_table(self) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            citations=ReportFixtures.make_citations_mock(total=4, matched=2)
        )

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._table_rows(),
            [
                ["Métrica", "Valor"],
                ["Calidad General", "8.0/10"],
                ["Gramática y Ortografía", "7.5/10"],
                ["Estructura", "✓ Válida"],
                ["Errores APA 7", "Sin errores"],
                ["Tasa de Coincidencia", "50.0%"],
            ],
        )
        table_style = self.document.tables[0].style
        assert table_style is not None
        self.assertEqual(table_style.name, self.settings.table_style)

    def test_render_marks_metrics_table_header_bold(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        header_runs = [
            run
            for cell in self.document.tables[0].rows[0].cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
        ]
        self.assertTrue(all(run.font.bold is True for run in header_runs))

    def test_render_counts_apa_errors_and_incomplete_structure(self) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            structure=ReportFixtures.make_structure_mock(is_valid=False),
            apa_validation=ReportFixtures.make_apa_validation_mock(violations=["a", "b"]),
        )

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(self._table_rows()[3], ["Estructura", "✗ Incompleta"])
        self.assertEqual(self._table_rows()[4], ["Errores APA 7", "2 detectados"])

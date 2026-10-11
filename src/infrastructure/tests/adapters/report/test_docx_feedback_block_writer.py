from unittest import TestCase

from docx import Document
from docx.shared import Inches

from src.domain.enums.feedback_block_kind import FeedbackBlockKind
from src.infrastructure.adapters.report.docx_feedback_block_payload import (
    DocxFeedbackBlockPayload,
)
from src.infrastructure.adapters.report.docx_feedback_block_writer import DocxFeedbackBlockWriter
from src.infrastructure.adapters.report.docx_markdown_run_writer import DocxMarkdownRunWriter
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxFeedbackBlockWriter(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.markdown_writer = DocxMarkdownRunWriter()
        self.writer = DocxFeedbackBlockWriter(
            markdown_run_writer=self.markdown_writer,
            docx_report_settings=self.settings,
        )
        self.document = Document()

    def test_write_blocks_renders_title_block_with_bold_run(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {"kind": "title", "text": "Fortalezas", "level": 0, "marker": ""}
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        self.assertEqual(len(self.document.paragraphs), 1)
        paragraph = self.document.paragraphs[0]
        self.assertEqual(paragraph.text, "Fortalezas")
        self.assertTrue(paragraph.runs[0].bold)

    def test_write_blocks_renders_level_zero_item_with_expected_indent(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento de nivel cero.",
                "level": 0,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[0]),
        )
        self.assertTrue(paragraph.text.startswith("• "))

    def test_write_blocks_renders_level_one_item_with_expected_indent(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento de nivel uno.",
                "level": 1,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[1]),
        )

    def test_write_blocks_renders_level_two_item_with_expected_indent(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento de nivel dos.",
                "level": 2,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[2]),
        )

    def test_write_blocks_clamps_indent_for_levels_exceeding_configured_levels(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento de nivel profundo.",
                "level": 5,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[2]),
        )

    def test_write_blocks_preserves_marker_already_ending_with_space(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Punto numerado.",
                "level": 0,
                "marker": "1. ",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertTrue(paragraph.text.startswith("1. Punto"))
        self.assertFalse(paragraph.text.startswith("1.  "))

    def test_write_blocks_renders_nested_markdown_in_item_text(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Subitem con frase **altamente relevante** destacada.",
                "level": 1,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertNotIn("**", paragraph.text)
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        self.assertIn("altamente relevante", bold_runs)

    def test_write_blocks_falls_back_to_markdown_paragraph_for_text_kind(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "text",
                "text": "Texto sin estructura especial pero con **énfasis**.",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertIsNone(paragraph.paragraph_format.left_indent)
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        self.assertIn("énfasis", bold_runs)

    def test_write_blocks_falls_back_to_text_writer_when_kind_is_unknown(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "unrecognized_kind",
                "text": "Texto con tipo desconocido y **formato**.",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertIsNone(paragraph.paragraph_format.left_indent)
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        self.assertIn("formato", bold_runs)

    def test_write_blocks_falls_back_to_text_writer_when_kind_is_missing(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "text": "Texto sin clave kind explícita.",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertIsNone(paragraph.paragraph_format.left_indent)
        self.assertEqual(paragraph.text, "Texto sin clave kind explícita.")

    def test_write_blocks_renders_title_when_kind_is_enum_value(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": FeedbackBlockKind.TITLE.value,
                "text": "Título con enum",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(paragraph.text, "Título con enum")
        self.assertTrue(paragraph.runs[0].bold)

    def test_write_blocks_renders_item_when_kind_is_enum_value(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": FeedbackBlockKind.ITEM.value,
                "text": "Elemento con enum",
                "level": 1,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[1]),
        )
        self.assertTrue(paragraph.text.startswith("• Elemento"))

    def test_write_blocks_renders_text_when_kind_is_enum_value(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": FeedbackBlockKind.TEXT.value,
                "text": "Párrafo de texto con enum",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertIsNone(paragraph.paragraph_format.left_indent)
        self.assertEqual(paragraph.text, "Párrafo de texto con enum")

    def test_write_blocks_clamps_indent_when_level_is_negative(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento con nivel negativo.",
                "level": -3,
                "marker": "•",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(
            paragraph.paragraph_format.left_indent,
            Inches(self.settings.feedback_item_indent_inches[0]),
        )

    def test_write_blocks_renders_empty_marker_without_prefix_spacing(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento sin marcador visible.",
                "level": 0,
                "marker": "",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertEqual(paragraph.text, "Elemento sin marcador visible.")

    def test_write_blocks_appends_space_to_marker_without_trailing_space(self) -> None:
        blocks: list[DocxFeedbackBlockPayload] = [
            {
                "kind": "item",
                "text": "Elemento con marcador sin espacio final.",
                "level": 0,
                "marker": "->",
            }
        ]

        self.writer.write_blocks(document=self.document, blocks=blocks)

        paragraph = self.document.paragraphs[0]
        self.assertTrue(paragraph.text.startswith("-> Elemento"))

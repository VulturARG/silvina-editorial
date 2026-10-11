from unittest import TestCase

from docx import Document

from src.domain.dtos.grammar_error_dto import GrammarErrorDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxGrammarAnalysisSection(TestCase):
    def setUp(self) -> None:
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def _render_errors(self, errors: list[GrammarErrorDTO], **setting_overrides) -> None:
        section = ReportFixtures.make_grammar_analysis_section(
            docx_report_settings=ReportFixtures.make_settings(**setting_overrides)
        )
        report_input = ReportFixtures.make_report_input_dto(
            grammar=ReportFixtures.make_grammar_mock(errors=errors)
        )
        section.render(document=self.document, report_input=report_input)

    def _make_error(
        self,
        message: str = "error",
        context: str = "ctx",
        replacements: list[str] | None = None,
    ) -> GrammarErrorDTO:
        return GrammarErrorDTO(
            number=1,
            message=message,
            context=context,
            offset=0,
            length=1,
            replacements=replacements if replacements is not None else [],
        )

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(ReportFixtures.make_grammar_analysis_section(), DocxReportSection)

    def test_render_reads_score_and_status_from_report_input_grammar(self) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            grammar=ReportFixtures.make_grammar_mock(score=6.5)
        )

        ReportFixtures.make_grammar_analysis_section().render(
            document=self.document, report_input=report_input
        )

        self.assertEqual(
            self._paragraph_texts(),
            ["📝 GRAMÁTICA Y ORTOGRAFÍA", "Puntuación: 6.5/10", "Estado: Good grammar"],
        )

    def test_render_lists_errors_under_bold_label(self) -> None:
        self._render_errors(
            [self._make_error(message="Falta tilde", context="cancion", replacements=["canción"])]
        )

        self.assertEqual(
            self._paragraph_texts()[3:],
            [
                "",
                "Errores Detectados:",
                "Falta tilde",
                '   Contexto: "cancion"',
                "   Sugerencia: canción",
            ],
        )
        self.assertIs(self.document.paragraphs[4].runs[0].bold, True)

    def test_render_omits_suggestion_when_error_has_no_replacements(self) -> None:
        self._render_errors([self._make_error()])

        self.assertEqual(self._paragraph_texts()[5:], ["error", '   Contexto: "ctx"'])

    def test_grammar_errors_limited_by_max_errors_displayed(self) -> None:
        errors = [self._make_error(message=f"error {index}") for index in range(3)]

        self._render_errors(errors, max_errors_displayed=1)

        error_paragraphs = [text for text in self._paragraph_texts() if text.startswith("error")]
        self.assertEqual(len(error_paragraphs), 1)

    def test_grammar_context_truncated_by_context_truncation_limit(self) -> None:
        self._render_errors([self._make_error(context="x" * 50)], context_truncation_limit=10)

        context_paragraphs = [text for text in self._paragraph_texts() if "Contexto" in text]
        self.assertEqual(context_paragraphs[0], '   Contexto: "' + "x" * 10 + '..."')

    def test_grammar_replacements_limited_by_max_replacements(self) -> None:
        self._render_errors([self._make_error(replacements=["a", "b", "c"])], max_replacements=1)

        suggestion_paragraphs = [text for text in self._paragraph_texts() if "Sugerencia" in text]
        self.assertEqual(suggestion_paragraphs[0], "   Sugerencia: a")

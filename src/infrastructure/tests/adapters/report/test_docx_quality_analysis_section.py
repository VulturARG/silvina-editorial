from unittest import TestCase

from docx import Document
from docx.shared import Pt, RGBColor

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxQualityAnalysisSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings(
            score_high_threshold=8.0, score_medium_threshold=6.0
        )
        self.section = ReportFixtures.make_quality_analysis_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def _render_with_dimension(self, dimension_data: dict) -> None:
        quality = ReportFixtures.make_quality_mock()
        quality.dimension_scores = {"claridad": dimension_data}
        report_input = ReportFixtures.make_report_input_dto(quality=quality)
        self.section.render(document=self.document, report_input=report_input)

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_heading_and_overall_score(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(
            self._paragraph_texts(),
            ["⭐ ANÁLISIS DE CALIDAD SEMÁNTICA", "Puntuación General: 8.0/10"],
        )
        label_run, score_run = self.document.paragraphs[1].runs
        self.assertIs(label_run.bold, True)
        self.assertIs(score_run.bold, True)
        self.assertEqual(score_run.font.size, Pt(self.settings.score_font_size_pt))

    def test_render_colors_overall_score_by_threshold(self) -> None:
        cases = (
            (8.0, self.settings.publishable_color_rgb),
            (6.0, self.settings.warning_color_rgb),
            (5.9, self.settings.reject_color_rgb),
        )
        for overall_score, expected_color in cases:
            with self.subTest(overall_score=overall_score):
                document = Document()
                report_input = ReportFixtures.make_report_input_dto(
                    quality=ReportFixtures.make_quality_mock(overall_score=overall_score)
                )

                self.section.render(document=document, report_input=report_input)

                score_run = document.paragraphs[1].runs[1]
                self.assertEqual(score_run.font.color.rgb, RGBColor(*expected_color))

    def test_render_adds_dimension_heading_score_and_plain_feedback(self) -> None:
        self._render_with_dimension({"score": 7.0, "feedback": "Texto claro."})

        self.assertEqual(
            self._paragraph_texts()[2:],
            ["Claridad", "Puntuación: 7.0/10", "Texto claro."],
        )
        dimension_style = self.document.paragraphs[2].style
        assert dimension_style is not None
        self.assertEqual(dimension_style.name, "Heading 3")

    def test_render_prefers_feedback_blocks_over_plain_feedback(self) -> None:
        self._render_with_dimension(
            {
                "score": 7.0,
                "feedback": "Texto ignorado.",
                "feedback_blocks": [{"kind": "title", "text": "Fortalezas"}],
            }
        )

        self.assertEqual(
            self._paragraph_texts()[2:],
            ["Claridad", "Puntuación: 7.0/10", "Fortalezas"],
        )
        self.assertIs(self.document.paragraphs[4].runs[0].bold, True)

    def test_render_omits_feedback_when_dimension_has_none(self) -> None:
        self._render_with_dimension({"score": 7.0})

        self.assertEqual(self._paragraph_texts()[2:], ["Claridad", "Puntuación: 7.0/10"])

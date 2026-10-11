from datetime import datetime
from unittest import TestCase

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures

INJECTED_TIME = datetime(2001, 2, 3, 4, 5)


class TestDocxFooterSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_footer_section(
            clock=lambda: INJECTED_TIME, docx_report_settings=self.settings
        )
        self.document = Document()
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_blank_paragraph_and_separator_line(self) -> None:
        self.assertEqual(self.document.paragraphs[0].text, "")
        self.assertEqual(self.document.paragraphs[1].text, "─" * 80)

    def test_render_adds_generated_by_line_with_injected_clock_timestamp(self) -> None:
        self.assertEqual(
            self.document.paragraphs[2].text,
            "Generado por Silvina Editorial Assistant v0.9 | 03/02/2001 04:05",
        )

    def test_render_styles_generated_by_line_as_centered_italic_metadata(self) -> None:
        footer_paragraph = self.document.paragraphs[2]

        self.assertEqual(footer_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(len(footer_paragraph.runs), 2)
        for run in footer_paragraph.runs:
            self.assertTrue(run.italic)
            self.assertEqual(run.font.size, Pt(self.settings.metadata_font_size_pt))
            self.assertEqual(run.font.color.rgb, RGBColor(*self.settings.neutral_color_rgb))

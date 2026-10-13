from unittest import TestCase

from docx import Document
from docx.shared import RGBColor

from src.domain.dtos.eumic_violation_dto import EumicViolationDTO
from src.domain.enums.eumic_category import EumicCategory
from src.domain.enums.severity_level import SeverityLevel
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures

HEADING = "📏 NORMAS EDITORIALES NO CUMPLIDAS (EUMIC)"
COMPLIANT_NOTICE = "✓ El documento cumple las normas editoriales EUMIC evaluadas"
WARNING_TEXT = (
    "No se reciben trabajos que incumplan estas normas editoriales. "
    "Corrija los puntos indicados antes de reenviar el manuscrito."
)


class TestDocxEumicViolationsSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_eumic_violations_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def _render(self, violations: list[EumicViolationDTO]) -> None:
        report_input = ReportFixtures.make_report_input_dto(eumic_violations=violations)
        self.section.render(document=self.document, report_input=report_input)

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_writes_level_one_heading(self) -> None:
        self._render([])

        self.assertEqual(self._paragraph_texts()[0], HEADING)
        heading_style = self.document.paragraphs[0].style
        assert heading_style is not None
        self.assertEqual(heading_style.name, "Heading 1")

    def test_render_reports_compliant_document_in_publishable_color(self) -> None:
        self._render([])

        self.assertEqual(self._paragraph_texts(), [HEADING, COMPLIANT_NOTICE])
        self.assertEqual(
            self.document.paragraphs[1].runs[0].font.color.rgb,
            RGBColor(*self.settings.publishable_color_rgb),
        )

    def test_render_omits_warning_text_when_document_is_compliant(self) -> None:
        self._render([])

        self.assertNotIn(WARNING_TEXT, self._paragraph_texts())

    def test_render_shows_warning_in_bold_reject_color_before_violations(self) -> None:
        self._render([self._violation(EumicCategory.FORMAT, "Margen superior")])

        self.assertEqual(self._paragraph_texts()[1], WARNING_TEXT)
        warning_run = self.document.paragraphs[1].runs[0]
        self.assertIs(warning_run.bold, True)
        self.assertEqual(warning_run.font.color.rgb, RGBColor(*self.settings.reject_color_rgb))

    def test_render_does_not_show_compliant_notice_when_violations_exist(self) -> None:
        self._render([self._violation(EumicCategory.FORMAT, "Margen superior")])

        self.assertNotIn(COMPLIANT_NOTICE, self._paragraph_texts())

    def test_render_groups_violations_by_category_in_first_seen_order(self) -> None:
        self._render(
            [
                self._violation(EumicCategory.FORMAT, "Margen superior"),
                self._violation(EumicCategory.FIGURES, "Figura sin fuente"),
                self._violation(EumicCategory.FORMAT, "Fuente no estándar"),
            ]
        )

        self.assertEqual(
            self._paragraph_texts()[2:],
            [
                "Formato General",
                "Advertencia: Margen superior",
                "Advertencia: Fuente no estándar",
                "Figuras",
                "Advertencia: Figura sin fuente",
            ],
        )

    def test_render_bolds_category_sub_label(self) -> None:
        self._render([self._violation(EumicCategory.TABLES, "Tabla sin título")])

        category_paragraph = self.document.paragraphs[2]
        self.assertEqual(category_paragraph.text, "Tablas")
        self.assertIs(category_paragraph.runs[0].bold, True)

    def test_render_lists_each_violation_as_list_bullet(self) -> None:
        self._render([self._violation(EumicCategory.FORMAT, "Margen superior")])

        bullet_style = self.document.paragraphs[3].style
        assert bullet_style is not None
        self.assertEqual(bullet_style.name, "List Bullet")

    def test_render_shows_every_severity_with_its_spanish_label(self) -> None:
        self._render(
            [
                self._violation(EumicCategory.FORMAT, "Info", SeverityLevel.INFO),
                self._violation(EumicCategory.FORMAT, "Warn", SeverityLevel.WARNING),
                self._violation(EumicCategory.FORMAT, "Err", SeverityLevel.ERROR),
                self._violation(EumicCategory.FORMAT, "Crit", SeverityLevel.CRITICAL),
            ]
        )

        self.assertEqual(
            self._paragraph_texts()[3:],
            [
                "Información: Info",
                "Advertencia: Warn",
                "Error: Err",
                "Crítico: Crit",
            ],
        )

    def test_render_puts_details_on_a_following_line_of_the_bullet(self) -> None:
        self._render(
            [
                self._violation(
                    EumicCategory.FORMAT,
                    "Margen superior",
                    details="Requerido: 2.5 cm, Actual: 1.00 cm",
                )
            ]
        )

        self.assertEqual(
            self._paragraph_texts()[3],
            "Advertencia: Margen superior\nRequerido: 2.5 cm, Actual: 1.00 cm",
        )

    def test_render_omits_line_break_when_details_are_empty(self) -> None:
        self._render([self._violation(EumicCategory.FORMAT, "Margen superior")])

        self.assertEqual(self._paragraph_texts()[3], "Advertencia: Margen superior")

    @staticmethod
    def _violation(
        category: EumicCategory,
        message: str,
        severity: SeverityLevel = SeverityLevel.WARNING,
        details: str = "",
    ) -> EumicViolationDTO:
        return EumicViolationDTO(
            category=category,
            message=message,
            severity=severity,
            details=details,
        )

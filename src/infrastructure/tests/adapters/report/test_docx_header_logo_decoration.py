from datetime import datetime
from os.path import join
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Pt, RGBColor

from src.infrastructure.adapters.report.docx_page_decoration import DocxPageDecoration
from src.infrastructure.resources.assets import ASSETS_DIR
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures

DECORATION_LOGGER_NAME = "src.infrastructure.adapters.report.docx_header_logo_decoration"
INJECTED_TIME = datetime(2001, 2, 3, 4, 5)
HIDDEN_BORDER_NAMES = ["top", "left", "bottom", "right", "insideH", "insideV"]
HIDDEN_BORDER_ATTRIBUTE_NAMES = ["w:val", "w:sz", "w:space", "w:color"]
HIDDEN_BORDER_ATTRIBUTE_VALUES = ["none", "0", "0", "auto"]


class TestDocxHeaderLogoDecoration(TestCase):
    def setUp(self) -> None:
        self._temporary_directory = TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self._missing_logo_path = join(self._temporary_directory.name, "missing_logo.jpg")
        self._readable_logo_path = join(ASSETS_DIR, "logo.jpg")
        self.settings = ReportFixtures.make_settings()
        self.document = Document()

    def _decoration(self, logo_path: str | None) -> DocxPageDecoration:
        return ReportFixtures.make_header_logo_decoration(
            logo_path=logo_path,
            clock=lambda: INJECTED_TIME,
            docx_report_settings=self.settings,
        )

    def _first_section(self) -> Section:
        return self.document.sections[0]

    def _header_image_count(self) -> int:
        header_relationships = self._first_section().header.part.rels.values()
        return len(
            [
                relationship
                for relationship in header_relationships
                if relationship.reltype == RELATIONSHIP_TYPE.IMAGE
            ]
        )

    def test_is_docx_page_decoration(self) -> None:
        self.assertIsInstance(self._decoration(logo_path=None), DocxPageDecoration)

    def test_apply_leaves_header_untouched_when_logo_path_is_none(self) -> None:
        self._decoration(logo_path=None).apply(document=self.document)

        self.assertTrue(self._first_section().header.is_linked_to_previous)

    def test_apply_adds_one_borderless_two_cell_header_table(self) -> None:
        self._decoration(logo_path=self._readable_logo_path).apply(document=self.document)

        header_tables = self._first_section().header.tables
        self.assertEqual(len(header_tables), 1)
        self.assertEqual(len(header_tables[0].rows), 1)
        self.assertEqual(len(header_tables[0].columns), 2)
        hidden_borders = header_tables[0]._element.xpath("./w:tblPr/w:tblBorders/*")
        self.assertEqual(
            [border.tag for border in hidden_borders],
            [qn(f"w:{border_name}") for border_name in HIDDEN_BORDER_NAMES],
        )
        for border in hidden_borders:
            self.assertEqual(
                [
                    border.get(qn(attribute_name))
                    for attribute_name in HIDDEN_BORDER_ATTRIBUTE_NAMES
                ],
                HIDDEN_BORDER_ATTRIBUTE_VALUES,
            )

    def test_apply_writes_generated_by_label_and_injected_clock_date_in_left_cell(self) -> None:
        self._decoration(logo_path=self._readable_logo_path).apply(document=self.document)

        left_cell = self._first_section().header.tables[0].rows[0].cells[0]
        self.assertEqual(
            left_cell.text,
            "Generado por Silvina Editorial Assistant v0.9\n"
            + INJECTED_TIME.strftime("%d de %B de %Y"),
        )

    def test_apply_styles_left_cell_runs_as_metadata_with_italic_label(self) -> None:
        self._decoration(logo_path=self._readable_logo_path).apply(document=self.document)

        label_run, date_run = (
            self._first_section().header.tables[0].rows[0].cells[0].paragraphs[0].runs
        )
        self.assertTrue(label_run.italic)
        self.assertIsNone(date_run.italic)
        for run in (label_run, date_run):
            self.assertEqual(run.font.size, Pt(self.settings.metadata_font_size_pt))
            self.assertEqual(run.font.color.rgb, RGBColor(*self.settings.neutral_color_rgb))

    def test_apply_adds_logo_picture_when_logo_file_is_readable(self) -> None:
        self._decoration(logo_path=self._readable_logo_path).apply(document=self.document)

        self.assertEqual(self._header_image_count(), 1)

    def test_apply_builds_header_table_without_picture_when_logo_file_is_missing(self) -> None:
        with self.assertLogs(DECORATION_LOGGER_NAME, level="WARNING"):
            self._decoration(logo_path=self._missing_logo_path).apply(document=self.document)

        header = self._first_section().header
        self.assertEqual(len(header.tables), 1)
        self.assertIn("Generado por", header.tables[0].rows[0].cells[0].text)
        self.assertEqual(self._header_image_count(), 0)

    def test_apply_warns_once_with_logo_path_when_logo_file_is_missing(self) -> None:
        with self.assertLogs(DECORATION_LOGGER_NAME, level="WARNING") as captured_logs:
            self._decoration(logo_path=self._missing_logo_path).apply(document=self.document)

        self.assertEqual(len(captured_logs.records), 1)
        self.assertIn(self._missing_logo_path, captured_logs.output[0])

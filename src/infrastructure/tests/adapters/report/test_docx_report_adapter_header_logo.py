from os.path import join
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from src.infrastructure.resources.assets import ASSETS_DIR
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures

ADAPTER_LOGGER_NAME = "src.infrastructure.adapters.report.docx_report_adapter"


class TestDocxReportAdapterHeaderLogo(TestCase):
    """Verify the header logo is skipped with a warning when the file is not readable."""

    def setUp(self) -> None:
        self._temporary_directory = TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self._output_path = join(self._temporary_directory.name, "report.docx")
        self._missing_logo_path = join(self._temporary_directory.name, "missing_logo.jpg")

    def test_export_succeeds_when_logo_file_is_missing(self) -> None:
        adapter = ReportFixtures.make_adapter(logo_path=self._missing_logo_path)

        with self.assertLogs(ADAPTER_LOGGER_NAME, level="WARNING"):
            export_succeeded = adapter.export(
                report_input=ReportFixtures.make_report_input_dto(),
                output_path=self._output_path,
            )

        self.assertTrue(export_succeeded)

    def test_header_table_is_built_without_picture_when_logo_file_is_missing(self) -> None:
        adapter = ReportFixtures.make_adapter(logo_path=self._missing_logo_path)

        with self.assertLogs(ADAPTER_LOGGER_NAME, level="WARNING"):
            adapter.export(
                report_input=ReportFixtures.make_report_input_dto(),
                output_path=self._output_path,
            )

        header = Document(self._output_path).sections[0].header
        self.assertEqual(len(header.tables), 1)
        self.assertIn("Generado por", header.tables[0].rows[0].cells[0].text)
        self.assertEqual(header._element.xpath(".//w:drawing"), [])

    def test_warning_mentions_logo_path_when_logo_file_is_missing(self) -> None:
        adapter = ReportFixtures.make_adapter(logo_path=self._missing_logo_path)

        with self.assertLogs(ADAPTER_LOGGER_NAME, level="WARNING") as captured_logs:
            adapter.export(
                report_input=ReportFixtures.make_report_input_dto(),
                output_path=self._output_path,
            )

        self.assertEqual(len(captured_logs.records), 1)
        self.assertIn(self._missing_logo_path, captured_logs.output[0])

    def test_header_contains_picture_when_logo_file_is_readable(self) -> None:
        adapter = ReportFixtures.make_adapter(logo_path=join(ASSETS_DIR, "logo.jpg"))

        adapter.export(
            report_input=ReportFixtures.make_report_input_dto(),
            output_path=self._output_path,
        )

        header = Document(self._output_path).sections[0].header
        self.assertEqual(len(header._element.xpath(".//w:drawing")), 1)

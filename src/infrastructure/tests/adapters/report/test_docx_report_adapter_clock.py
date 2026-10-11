from datetime import datetime
from os.path import join
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from src.infrastructure.resources.assets import ASSETS_DIR
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures

INJECTED_TIME = datetime(2001, 2, 3, 4, 5)


class TestDocxReportAdapterClock(TestCase):
    """Verify header and footer dates come from the injected clock."""

    def setUp(self) -> None:
        self._temporary_directory = TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self._output_path = join(self._temporary_directory.name, "report.docx")
        adapter = ReportFixtures.make_adapter(
            logo_path=join(ASSETS_DIR, "logo.jpg"),
            clock=lambda: INJECTED_TIME,
        )
        adapter.export(
            report_input=ReportFixtures.make_report_input_dto(),
            output_path=self._output_path,
        )
        self._document = Document(self._output_path)

    def test_header_date_comes_from_injected_clock(self) -> None:
        header_cell_text = self._document.sections[0].header.tables[0].rows[0].cells[0].text

        self.assertTrue(header_cell_text.endswith(INJECTED_TIME.strftime("%d de %B de %Y")))

    def test_footer_date_comes_from_injected_clock(self) -> None:
        footer_paragraphs = [
            paragraph.text
            for paragraph in self._document.paragraphs
            if paragraph.text.startswith("Generado por")
        ]

        self.assertEqual(len(footer_paragraphs), 1)
        self.assertTrue(footer_paragraphs[0].endswith("| 03/02/2001 04:05"))

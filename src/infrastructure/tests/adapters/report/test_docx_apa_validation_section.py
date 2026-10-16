from unittest import TestCase

from docx import Document
from docx.shared import RGBColor

from src.domain.dtos.apa_violation_dto import ApaViolationDTO
from src.domain.enums.apa_error_type import ApaErrorType
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxApaValidationSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_apa_validation_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _paragraph_texts(self) -> list[str]:
        return [paragraph.text for paragraph in self.document.paragraphs]

    def _make_violation(
        self, citation_text: str = "(Smith, 2020)", correction: str = "(Smith, 2020)"
    ) -> ApaViolationDTO:
        return ApaViolationDTO(
            citation_text=citation_text,
            error_type=ApaErrorType.CONJUNCTION_ERROR,
            location=5,
            explanation="Missing year",
            correction=correction,
        )

    def _render_violations(self, violations: list[ApaViolationDTO]) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            apa_validation=ReportFixtures.make_apa_validation_mock(violations=violations)
        )
        self.section.render(document=self.document, report_input=report_input)

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_reports_no_errors_in_publishable_color_when_citations_exist(self) -> None:
        self._render_violations([])

        self.assertEqual(
            self._paragraph_texts(),
            [
                "📖 VALIDACIÓN APA 7 (ESPAÑOL)",
                "✅ Sin errores de formato APA 7 detectados",
                "",
            ],
        )
        self.assertEqual(
            self.document.paragraphs[1].runs[0].font.color.rgb,
            RGBColor(*self.settings.publishable_color_rgb),
        )

    def test_render_reports_unvalidated_format_in_neutral_color_without_citations(self) -> None:
        report_input = ReportFixtures.make_report_input_dto(
            citations=ReportFixtures.make_citations_mock(total=0)
        )

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._paragraph_texts()[1],
            "ℹ️ Sin citaciones en texto detectadas — no se pudo validar formato APA 7",
        )
        self.assertEqual(
            self.document.paragraphs[1].runs[0].font.color.rgb,
            RGBColor(*self.settings.neutral_color_rgb),
        )

    def test_render_reads_violation_attributes(self) -> None:
        self._render_violations([self._make_violation()])

        self.assertEqual(
            self._paragraph_texts()[1:],
            [
                "CITACIÓN INCORRECTA (1):",
                "1. Citación: (Smith, 2020)",
                '   Ubicación: "5"',
                "   Problema: Missing year",
                "   Corrección: (Smith, 2020)",
            ],
        )
        group_style = self.document.paragraphs[1].style
        assert group_style is not None
        self.assertEqual(group_style.name, "Heading 3")

    def test_render_omits_correction_when_violation_has_none(self) -> None:
        self._render_violations([self._make_violation(correction="")])

        self.assertNotIn("Corrección", "\n".join(self._paragraph_texts()))

    def test_apa_violations_are_not_limited_by_max_errors_displayed(self) -> None:
        section = ReportFixtures.make_apa_validation_section(
            docx_report_settings=ReportFixtures.make_settings(max_errors_displayed=1)
        )
        violations = [self._make_violation(citation_text=f"cite{index}") for index in range(3)]
        report_input = ReportFixtures.make_report_input_dto(
            apa_validation=ReportFixtures.make_apa_validation_mock(violations=violations)
        )

        section.render(document=self.document, report_input=report_input)

        citation_paragraphs = [text for text in self._paragraph_texts() if "Citación:" in text]
        self.assertEqual(
            citation_paragraphs,
            ["1. Citación: cite0", "2. Citación: cite1", "3. Citación: cite2"],
        )
        self.assertEqual(self._paragraph_texts()[1], "CITACIÓN INCORRECTA (3):")

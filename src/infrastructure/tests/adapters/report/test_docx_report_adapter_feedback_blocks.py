from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from src.domain.dtos.quality_result_dto import QualityResultDTO
from src.domain.enums.quality_level import QualityLevel
from src.infrastructure.adapters.report.docx_report_adapter import DocxReportAdapter
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxReportAdapterFeedbackBlocks(TestCase):
    """Integration tests for rendering structured feedback blocks in Word documents."""

    def setUp(self) -> None:
        self._temporary_directory = TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self.output_path = str(Path(self._temporary_directory.name) / "report.docx")
        self.adapter = DocxReportAdapter(
            logo_path=None, docx_report_settings=ReportFixtures.make_settings()
        )

    def test_export_renders_structured_feedback_blocks_and_preserves_flat_dimensions(
        self,
    ) -> None:
        quality = QualityResultDTO(
            overall_score=8.5,
            quality_level=QualityLevel.GOOD,
            dimension_scores={
                "claridad": {
                    "score": 9.0,
                    "feedback": "Fortalezas: Primer punto. Segundo punto. 1. Punto numerado.",
                    "feedback_blocks": [
                        {
                            "kind": "title",
                            "text": "Fortalezas",
                            "level": 0,
                            "marker": "",
                        },
                        {
                            "kind": "item",
                            "text": "Primer aspecto positivo identificado.",
                            "level": 0,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Segundo aspecto positivo detallado.",
                            "level": 0,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Punto principal estructurado.",
                            "level": 0,
                            "marker": "1.",
                        },
                        {
                            "kind": "item",
                            "text": "Subitem con frase **altamente relevante** destacada.",
                            "level": 1,
                            "marker": "•",
                        },
                    ],
                },
                "coherencia": {
                    "score": 8.0,
                    "feedback": "Coherencia global satisfactoria sin bloques.",
                },
            },
        )
        report_input = ReportFixtures.make_report_input_dto(quality=quality)

        self.adapter.export(report_input=report_input, output_path=self.output_path)

        document = Document(self.output_path)
        paragraphs = document.paragraphs

        clarity_heading_index = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip() == "Claridad"
        )
        coherence_heading_index = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip() == "Coherencia"
        )

        clarity_paragraphs = paragraphs[clarity_heading_index + 1 : coherence_heading_index]
        coherence_paragraphs = paragraphs[coherence_heading_index + 1 : coherence_heading_index + 3]

        self.assertEqual(clarity_paragraphs[0].text, "Puntuación: 9.0/10")

        title_paragraph = clarity_paragraphs[1]
        self.assertEqual(title_paragraph.text, "Fortalezas")
        self.assertTrue(any(run.bold for run in title_paragraph.runs))

        first_item_paragraph = clarity_paragraphs[2]
        self.assertTrue(first_item_paragraph.text.startswith("•"))
        self.assertIn("Primer aspecto positivo identificado.", first_item_paragraph.text)

        second_item_paragraph = clarity_paragraphs[3]
        self.assertTrue(second_item_paragraph.text.startswith("•"))
        self.assertIn("Segundo aspecto positivo detallado.", second_item_paragraph.text)

        numbered_item_paragraph = clarity_paragraphs[4]
        self.assertTrue(numbered_item_paragraph.text.startswith("1."))
        self.assertIn("Punto principal estructurado.", numbered_item_paragraph.text)

        nested_item_paragraph = clarity_paragraphs[5]
        self.assertTrue(nested_item_paragraph.text.startswith("•"))
        self.assertIn("altamente relevante", nested_item_paragraph.text)
        self.assertNotIn("**", nested_item_paragraph.text)

        bold_runs = [run.text for run in nested_item_paragraph.runs if run.bold]
        self.assertIn("altamente relevante", bold_runs)

        nested_indent = nested_item_paragraph.paragraph_format.left_indent
        first_indent = first_item_paragraph.paragraph_format.left_indent
        self.assertIsNotNone(nested_indent)
        self.assertIsNotNone(first_indent)
        assert nested_indent is not None
        assert first_indent is not None
        self.assertGreater(
            nested_indent,
            first_indent,
        )

        self.assertEqual(coherence_paragraphs[0].text, "Puntuación: 8.0/10")
        self.assertEqual(
            coherence_paragraphs[1].text,
            "Coherencia global satisfactoria sin bloques.",
        )

    def test_export_renders_level_two_block_with_larger_indent_than_level_one(
        self,
    ) -> None:
        quality = QualityResultDTO(
            overall_score=8.5,
            quality_level=QualityLevel.GOOD,
            dimension_scores={
                "claridad": {
                    "score": 9.0,
                    "feedback": "Texto plano de respaldo.",
                    "feedback_blocks": [
                        {
                            "kind": "item",
                            "text": "Elemento de nivel cero.",
                            "level": 0,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Elemento hijo de nivel uno.",
                            "level": 1,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Elemento sub-hijo de nivel dos.",
                            "level": 2,
                            "marker": "•",
                        },
                    ],
                },
            },
        )
        report_input = ReportFixtures.make_report_input_dto(quality=quality)

        self.adapter.export(report_input=report_input, output_path=self.output_path)

        document = Document(self.output_path)
        paragraphs = document.paragraphs

        clarity_heading_index = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip() == "Claridad"
        )

        level_zero_paragraph = paragraphs[clarity_heading_index + 2]
        level_one_paragraph = paragraphs[clarity_heading_index + 3]
        level_two_paragraph = paragraphs[clarity_heading_index + 4]

        indent_zero = level_zero_paragraph.paragraph_format.left_indent
        indent_one = level_one_paragraph.paragraph_format.left_indent
        indent_two = level_two_paragraph.paragraph_format.left_indent

        self.assertIsNotNone(indent_zero)
        self.assertIsNotNone(indent_one)
        self.assertIsNotNone(indent_two)
        assert indent_zero is not None
        assert indent_one is not None
        assert indent_two is not None
        self.assertGreater(indent_one, indent_zero)
        self.assertGreater(indent_two, indent_one)

    def test_export_renders_italic_and_bold_runs_without_literal_asterisks(
        self,
    ) -> None:
        quality = QualityResultDTO(
            overall_score=8.5,
            quality_level=QualityLevel.GOOD,
            dimension_scores={
                "claridad": {
                    "score": 9.0,
                    "feedback": "Texto plano.",
                    "feedback_blocks": [
                        {
                            "kind": "item",
                            "text": "El texto emplea *correctamente* el *mecanismo* del *cual* depende.",
                            "level": 0,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Texto con **negrita** y *cursiva*, más **negrita con *cursiva anidada***.",
                            "level": 0,
                            "marker": "•",
                        },
                        {
                            "kind": "item",
                            "text": "Operación a * b, lista '* item', escala 1 x 10^11, nota* y *unpaired.",
                            "level": 0,
                            "marker": "•",
                        },
                    ],
                },
            },
        )
        report_input = ReportFixtures.make_report_input_dto(quality=quality)

        self.adapter.export(report_input=report_input, output_path=self.output_path)

        document = Document(self.output_path)
        paragraphs = document.paragraphs

        clarity_heading_index = next(
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip() == "Claridad"
        )

        emphasis_paragraph = paragraphs[clarity_heading_index + 2]
        mixed_paragraph = paragraphs[clarity_heading_index + 3]
        lone_paragraph = paragraphs[clarity_heading_index + 4]

        self.assertNotIn("*", emphasis_paragraph.text)
        italic_runs_in_first = [run.text for run in emphasis_paragraph.runs if run.italic]
        self.assertEqual(
            italic_runs_in_first,
            ["correctamente", "mecanismo", "cual"],
        )

        self.assertNotIn("*", mixed_paragraph.text)
        bold_runs = [run.text for run in mixed_paragraph.runs if run.bold]
        self.assertIn("negrita", bold_runs)
        self.assertIn("cursiva anidada", bold_runs)
        italic_runs_in_second = [run.text for run in mixed_paragraph.runs if run.italic]
        self.assertIn("cursiva", italic_runs_in_second)
        self.assertIn("cursiva anidada", italic_runs_in_second)

        self.assertIn("a * b", lone_paragraph.text)
        self.assertIn("'* item'", lone_paragraph.text)
        self.assertIn("1 x 10^11", lone_paragraph.text)
        self.assertIn("nota*", lone_paragraph.text)
        self.assertIn("*unpaired", lone_paragraph.text)

    def test_export_renders_dimension_heading_with_spanish_label_accent_for_argumentacion(
        self,
    ) -> None:
        quality = QualityResultDTO(
            overall_score=8.0,
            quality_level=QualityLevel.GOOD,
            dimension_scores={
                "argumentacion": {
                    "score": 8.0,
                    "feedback": "Texto de argumentación.",
                },
            },
        )
        report_input = ReportFixtures.make_report_input_dto(quality=quality)

        self.adapter.export(report_input=report_input, output_path=self.output_path)

        document = Document(self.output_path)
        paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs]

        self.assertIn("Argumentación", paragraph_texts)
        self.assertNotIn("Argumentacion", paragraph_texts)

from unittest import TestCase

from docx import Document

from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.tests.adapters.report.fixtures import ReportFixtures


class TestDocxCitationsAnalysisSection(TestCase):
    def setUp(self) -> None:
        self.settings = ReportFixtures.make_settings()
        self.section = ReportFixtures.make_citations_analysis_section(
            docx_report_settings=self.settings
        )
        self.document = Document()

    def _table_rows(self) -> list[list[str]]:
        return [[cell.text for cell in row.cells] for row in self.document.tables[0].rows]

    def test_is_docx_report_section(self) -> None:
        self.assertIsInstance(self.section, DocxReportSection)

    def test_render_adds_section_heading(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        self.assertEqual(self.document.paragraphs[0].text, "📚 ANÁLISIS DE CITAS Y REFERENCIAS")

    def test_render_adds_citation_metrics_table(self) -> None:
        citations = ReportFixtures.make_citations_mock(total=8, matched=6, references=7)
        report_input = ReportFixtures.make_report_input_dto(citations=citations)

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(
            self._table_rows(),
            [
                ["Total de citas en texto", "8"],
                ["Total de referencias bibliográficas", "7"],
                ["Citas con referencia", "6"],
                ["Tasa de coincidencia", "75.0%"],
            ],
        )

    def test_render_applies_table_style_from_settings(self) -> None:
        self.section.render(
            document=self.document, report_input=ReportFixtures.make_report_input_dto()
        )

        table_style = self.document.tables[0].style
        assert table_style is not None
        self.assertEqual(table_style.name, self.settings.table_style)

    def test_render_shows_not_available_match_rate_without_citations(self) -> None:
        citations = ReportFixtures.make_citations_mock(total=0, matched=0, references=0)
        report_input = ReportFixtures.make_report_input_dto(citations=citations)

        self.section.render(document=self.document, report_input=report_input)

        self.assertEqual(self._table_rows()[3], ["Tasa de coincidencia", "N/A"])

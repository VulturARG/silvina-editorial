from collections.abc import Callable
from datetime import datetime

from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings

FOOTER_SEPARATOR = "─" * 66
FOOTER_TIMESTAMP_FORMAT = "%d/%m/%Y %H:%M"


class DocxFooterSection(DocxReportSection):
    """Render the closing separator and the generation stamp of the report."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        clock: Callable[[], datetime],
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._clock = clock

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the separator line and the application name, version and timestamp."""
        document.add_paragraph()
        document.add_paragraph(FOOTER_SEPARATOR)

        footer_paragraph = document.add_paragraph()
        footer_paragraph.add_run(
            f"Generado por {self._docx_report_settings.app_name} v{self._docx_report_settings.app_version} | "
        ).italic = True
        footer_paragraph.add_run(self._clock().strftime(FOOTER_TIMESTAMP_FORMAT)).italic = True
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer_paragraph.runs:
            run.font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
            run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

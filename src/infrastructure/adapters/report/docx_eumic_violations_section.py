from docx.document import Document as DocxDocument
from docx.shared import RGBColor

from src.domain.dtos.eumic_violation_dto import EumicViolationDTO
from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)


class DocxEumicViolationsSection(DocxReportSection):
    """Render the unmet EUMIC editorial norms grouped by category."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the compliance notice, or the rejection warning followed by the violations."""
        self._section_heading_writer.write(
            document=document, text="📏 NORMAS EDITORIALES NO CUMPLIDAS (EUMIC)", level=1
        )

        violations = report_input.eumic_violations

        if not violations:
            document.add_paragraph().add_run(
                "✓ El documento cumple las normas editoriales EUMIC evaluadas"
            ).font.color.rgb = RGBColor(*self._docx_report_settings.publishable_color_rgb)
            return

        warning_run = document.add_paragraph().add_run(
            "Sección informativa: estas observaciones no modifican el veredicto de publicación del análisis. Las normas editoriales EUMIC no admiten trabajos que las incumplan; corrija los puntos indicados antes de reenviar el manuscrito."
        )
        warning_run.bold = True
        warning_run.font.color.rgb = RGBColor(*self._docx_report_settings.reject_color_rgb)

        for category, category_violations in self._group_by_category(violations).items():
            document.add_paragraph().add_run(category).bold = True
            for violation in category_violations:
                self._write_violation(document=document, violation=violation)

    def _group_by_category(
        self, violations: list[EumicViolationDTO]
    ) -> dict[str, list[EumicViolationDTO]]:
        grouped_violations: dict[str, list[EumicViolationDTO]] = {}
        for violation in violations:
            category_label = self._category_label(violation.category)
            grouped_violations.setdefault(category_label, []).append(violation)
        return grouped_violations

    @staticmethod
    def _category_label(category: str) -> str:
        return getattr(category, "value", category)

    @staticmethod
    def _write_violation(document: DocxDocument, violation: EumicViolationDTO) -> None:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{violation.severity.label}: ").bold = True
        paragraph.add_run(violation.message)
        if violation.details:
            paragraph.add_run().add_break()
            paragraph.add_run(violation.details)

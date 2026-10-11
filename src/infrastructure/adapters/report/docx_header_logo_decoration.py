from collections.abc import Callable
from datetime import datetime
from logging import getLogger
from os import R_OK, access
from os.path import isfile

from docx.document import Document as DocxDocument
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from src.infrastructure.adapters.report.docx_page_decoration import DocxPageDecoration
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings

HEADER_TABLE_ROW_COUNT = 1
HEADER_TABLE_COLUMN_COUNT = 2
HEADER_ROW_INDEX = 0
LEFT_CELL_INDEX = 0
RIGHT_CELL_INDEX = 1
HEADER_DATE_FORMAT = "%d de %B de %Y"
HIDDEN_TABLE_BORDER_NAMES = ("top", "left", "bottom", "right", "insideH", "insideV")
HIDDEN_BORDER_ATTRIBUTES = (
    ("w:val", "none"),
    ("w:sz", "0"),
    ("w:space", "0"),
    ("w:color", "auto"),
)

logger = getLogger(__name__)


class DocxHeaderLogoDecoration(DocxPageDecoration):
    """Add the borderless header table with the generation stamp and the logo picture."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        clock: Callable[[], datetime],
        logo_path: str | None,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._clock = clock
        self._logo_path = logo_path

    def apply(self, document: DocxDocument) -> None:
        """Build the header table when a logo path is configured, otherwise leave the header."""
        if self._logo_path is None:
            return
        header_table = self._add_header_table(document=document)
        self._fill_left_cell(header_table=header_table)
        self._fill_right_cell(header_table=header_table, logo_path=self._logo_path)
        self._hide_table_borders(header_table=header_table)

    def _add_header_table(self, document: DocxDocument) -> Table:
        header = document.sections[0].header
        header.paragraphs[0].clear()
        return header.add_table(
            rows=HEADER_TABLE_ROW_COUNT,
            cols=HEADER_TABLE_COLUMN_COUNT,
            width=Inches(self._docx_report_settings.header_table_width_inches),
        )

    def _fill_left_cell(self, header_table: Table) -> None:
        left_cell = header_table.cell(HEADER_ROW_INDEX, LEFT_CELL_INDEX)
        left_cell.width = Inches(self._docx_report_settings.header_left_cell_width_inches)
        left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        left_paragraph = left_cell.paragraphs[0]
        left_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        app_label_run = left_paragraph.add_run(
            f"Generado por {self._docx_report_settings.app_name} v{self._docx_report_settings.app_version}\n"
        )
        app_label_run.italic = True
        self._style_metadata_run(run=app_label_run)

        date_run = left_paragraph.add_run(self._clock().strftime(HEADER_DATE_FORMAT))
        self._style_metadata_run(run=date_run)

    def _fill_right_cell(self, header_table: Table, logo_path: str) -> None:
        right_cell = header_table.cell(HEADER_ROW_INDEX, RIGHT_CELL_INDEX)
        right_cell.width = Inches(self._docx_report_settings.header_right_cell_width_inches)
        right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right_paragraph = right_cell.paragraphs[0]
        right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._add_logo_picture(paragraph=right_paragraph, logo_path=logo_path)

    def _add_logo_picture(self, paragraph: Paragraph, logo_path: str) -> None:
        if not (isfile(logo_path) and access(logo_path, R_OK)):
            logger.warning("Header logo is not readable, skipping picture: %s", logo_path)
            return
        logo_run = paragraph.add_run()
        logo_run.add_picture(logo_path, width=Inches(self._docx_report_settings.logo_width_inches))

    def _style_metadata_run(self, run: Run) -> None:
        run.font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
        run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

    def _hide_table_borders(self, header_table: Table) -> None:
        table_properties = self._table_properties(header_table=header_table)
        table_borders = OxmlElement("w:tblBorders")
        for border_name in HIDDEN_TABLE_BORDER_NAMES:
            table_borders.append(self._hidden_border(border_name=border_name))
        table_properties.append(table_borders)

    def _table_properties(self, header_table: Table) -> BaseOxmlElement:
        table_element = header_table._element
        table_properties = table_element.tblPr
        if table_properties is None:
            table_properties = OxmlElement("w:tblPr")
            table_element.insert(0, table_properties)
        return table_properties

    def _hidden_border(self, border_name: str) -> BaseOxmlElement:
        border = OxmlElement(f"w:{border_name}")
        for attribute_name, attribute_value in HIDDEN_BORDER_ATTRIBUTES:
            border.set(qn(attribute_name), attribute_value)
        return border

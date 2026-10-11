from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph


class DocxLabeledParagraphWriter:
    """Render labeled paragraphs with a bold label followed by a plain value."""

    def write(self, document: DocxDocument, label: str, value: str) -> Paragraph:
        """Add a paragraph containing a bold label run and a plain value run."""
        paragraph = document.add_paragraph()
        paragraph.add_run(label).bold = True
        paragraph.add_run(value)
        return paragraph

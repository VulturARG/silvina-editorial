from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from src.infrastructure.adapters.report.docx_page_decoration import DocxPageDecoration
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings

PAGE_FIELD_INSTRUCTION = "PAGE"
TOTAL_PAGES_FIELD_INSTRUCTION = "NUMPAGES"
PAGE_NUMBER_SEPARATOR = " de "
FIELD_BEGIN_CHARACTER_TYPE = "begin"
FIELD_END_CHARACTER_TYPE = "end"
PRESERVED_SPACE = "preserve"


class DocxPageNumberDecoration(DocxPageDecoration):
    """Add the centered "PAGE de NUMPAGES" page counter to the footer."""

    def __init__(self, docx_report_settings: DocxReportSettings) -> None:
        self._docx_report_settings = docx_report_settings

    def apply(self, document: DocxDocument) -> None:
        """Replace the footer content with the current page and total pages fields."""
        footer_paragraph = document.sections[0].footer.paragraphs[0]
        footer_paragraph.clear()
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        self._add_field_run(paragraph=footer_paragraph, field_instruction=PAGE_FIELD_INSTRUCTION)
        footer_paragraph.add_run(PAGE_NUMBER_SEPARATOR)
        self._add_field_run(
            paragraph=footer_paragraph, field_instruction=TOTAL_PAGES_FIELD_INSTRUCTION
        )

        for footer_run in footer_paragraph.runs:
            footer_run.font.size = Pt(self._docx_report_settings.page_number_font_size_pt)
            footer_run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

    def _add_field_run(self, paragraph: Paragraph, field_instruction: str) -> None:
        field_run_element = paragraph.add_run()._element
        field_run_element.append(self._field_character(FIELD_BEGIN_CHARACTER_TYPE))
        field_run_element.append(self._field_instruction(field_instruction))
        field_run_element.append(self._field_character(FIELD_END_CHARACTER_TYPE))

    def _field_character(self, field_character_type: str) -> BaseOxmlElement:
        field_character = OxmlElement("w:fldChar")
        field_character.set(qn("w:fldCharType"), field_character_type)
        return field_character

    def _field_instruction(self, field_instruction: str) -> BaseOxmlElement:
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), PRESERVED_SPACE)
        instruction.text = field_instruction
        return instruction

from docx.document import Document as DocxDocument
from docx.shared import RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)


class DocxStructureValidationSection(DocxReportSection):
    """Render the EUMIC structure validation result and its missing sections."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the structure verdict and list missing sections when it is incomplete."""
        self._section_heading_writer.write(
            document=document, text="📋 VALIDACIÓN DE ESTRUCTURA (EUMIC)", level=1
        )

        structure = report_input.structure

        if structure.is_valid:
            document.add_paragraph().add_run(
                "✓ Estructura válida según normas EUMIC"
            ).font.color.rgb = RGBColor(*self._docx_report_settings.publishable_color_rgb)
            return

        document.add_paragraph().add_run("✗ Estructura incompleta").font.color.rgb = RGBColor(
            *self._docx_report_settings.reject_color_rgb
        )

        if not structure.missing_sections:
            return

        document.add_paragraph()
        document.add_paragraph().add_run("Secciones faltantes:").bold = True
        for missing_section in structure.missing_sections:
            document.add_paragraph(f"• {missing_section}", style="List Bullet")

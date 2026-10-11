from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.domain.enums.quality_dimension import QualityDimension
from src.infrastructure.adapters.report.docx_feedback_block_writer import (
    DocxFeedbackBlockWriter,
)
from src.infrastructure.adapters.report.docx_labeled_paragraph_writer import (
    DocxLabeledParagraphWriter,
)
from src.infrastructure.adapters.report.docx_markdown_run_writer import (
    DocxMarkdownRunWriter,
)
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)

DIMENSION_HEADING_LEVEL = 3


class DocxQualityAnalysisSection(DocxReportSection):
    """Render the overall semantic quality score and the per-dimension feedback."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
        labeled_paragraph_writer: DocxLabeledParagraphWriter,
        markdown_run_writer: DocxMarkdownRunWriter,
        feedback_block_writer: DocxFeedbackBlockWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer
        self._labeled_paragraph_writer = labeled_paragraph_writer
        self._markdown_run_writer = markdown_run_writer
        self._feedback_block_writer = feedback_block_writer

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the score colored by threshold and each dimension score with its feedback."""
        self._section_heading_writer.write(
            document=document, text="⭐ ANÁLISIS DE CALIDAD SEMÁNTICA", level=1
        )

        quality = report_input.quality

        paragraph = document.add_paragraph()
        paragraph.add_run("Puntuación General: ").bold = True
        score_run = paragraph.add_run(f"{quality.overall_score:.1f}/10")
        score_run.bold = True
        score_run.font.size = Pt(self._docx_report_settings.score_font_size_pt)
        score_run.font.color.rgb = RGBColor(*self._color_for_score(score=quality.overall_score))

        if quality.dimension_scores:
            for dimension_name, dimension_data in quality.dimension_scores.items():
                document.add_heading(
                    QualityDimension.label_for(dimension_name), level=DIMENSION_HEADING_LEVEL
                )

                self._labeled_paragraph_writer.write(
                    document=document,
                    label="Puntuación: ",
                    value=f"{dimension_data['score']:.1f}/10",
                )

                feedback_blocks = dimension_data.get("feedback_blocks")
                if feedback_blocks:
                    self._feedback_block_writer.write_blocks(
                        document=document, blocks=feedback_blocks
                    )
                elif dimension_data.get("feedback"):
                    self._markdown_run_writer.write_paragraph(
                        document=document,
                        text=dimension_data["feedback"],
                    )

    def _color_for_score(self, score: float) -> tuple[int, int, int]:
        if score >= self._docx_report_settings.score_high_threshold:
            return self._docx_report_settings.publishable_color_rgb
        if score >= self._docx_report_settings.score_medium_threshold:
            return self._docx_report_settings.warning_color_rgb
        return self._docx_report_settings.reject_color_rgb

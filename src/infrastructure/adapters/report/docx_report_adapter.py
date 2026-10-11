from __future__ import annotations

from collections import defaultdict
from collections.abc import Callable
from datetime import datetime
from logging import getLogger
from os import R_OK, access
from os.path import isfile

try:
    from docx import Document
    from docx.document import Document as DocxDocument
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.text.paragraph import Paragraph

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.domain.enums.publication_verdict import PublicationVerdict
from src.domain.enums.quality_dimension import QualityDimension
from src.domain.enums.recommendation_priority import RecommendationPriority
from src.domain.exceptions.report_errors import ReportExportUnavailable
from src.domain.report.report_export_port import ReportExportPort
from src.infrastructure.adapters.report.citation_match_rate_formatter import (
    CitationMatchRateFormatter,
)
from src.infrastructure.adapters.report.docx_feedback_block_writer import (
    DocxFeedbackBlockWriter,
)
from src.infrastructure.adapters.report.docx_labeled_paragraph_writer import (
    DocxLabeledParagraphWriter,
)
from src.infrastructure.adapters.report.docx_markdown_run_writer import (
    DocxMarkdownRunWriter,
)
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)

logger = getLogger(__name__)


class DocxReportAdapter(ReportExportPort):
    """Export analysis reports to professionally formatted Word (.docx) documents."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        markdown_run_writer: DocxMarkdownRunWriter,
        feedback_block_writer: DocxFeedbackBlockWriter,
        section_heading_writer: DocxSectionHeadingWriter,
        labeled_paragraph_writer: DocxLabeledParagraphWriter,
        citation_match_rate_formatter: CitationMatchRateFormatter,
        clock: Callable[[], datetime],
        logo_path: str | None = None,
    ) -> None:
        if not DOCX_AVAILABLE:
            raise ReportExportUnavailable()
        self._logo_path = logo_path
        self._docx_report_settings = docx_report_settings
        self._markdown_run_writer = markdown_run_writer
        self._feedback_block_writer = feedback_block_writer
        self._section_heading_writer = section_heading_writer
        self._labeled_paragraph_writer = labeled_paragraph_writer
        self._citation_match_rate_formatter = citation_match_rate_formatter
        self._clock = clock

    def _color_for_score(self, score: float) -> tuple[int, int, int]:
        if score >= self._docx_report_settings.score_high_threshold:
            return self._docx_report_settings.publishable_color_rgb
        if score >= self._docx_report_settings.score_medium_threshold:
            return self._docx_report_settings.warning_color_rgb
        return self._docx_report_settings.reject_color_rgb

    def export(self, report_input: ReportInputDTO, output_path: str) -> bool:
        """Write the report as a .docx file at output_path and return True on success."""
        document = Document()

        style = document.styles["Normal"]
        style.font.name = self._docx_report_settings.font_name
        style.font.size = Pt(self._docx_report_settings.base_font_size_pt)
        style.paragraph_format.line_spacing = self._docx_report_settings.line_spacing
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.space_before = Pt(0)

        self._add_header_logo(document=document)
        self._add_page_numbers(document=document)
        self._add_title_page(document=document, report_input=report_input)
        self._add_executive_summary(document=document, report_input=report_input)
        self._add_document_info(document=document, report_input=report_input)
        self._add_classification(document=document, report_input=report_input)
        self._add_quality_analysis(document=document, report_input=report_input)
        self._add_editorial_suitability(document=document, report_input=report_input)
        self._add_grammar_analysis(document=document, report_input=report_input)
        self._add_apa_validation(document=document, report_input=report_input)
        self._add_structure_validation(document=document, report_input=report_input)
        self._add_citations_analysis(document=document, report_input=report_input)
        self._add_recommendations(document=document, report_input=report_input)
        self._add_footer(document=document)

        document.save(output_path)
        return True

    def _add_title_page(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        title = document.add_heading("INFORME DE ANÁLISIS EDITORIAL", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.size = Pt(self._docx_report_settings.title_font_size_pt)
            run.font.color.rgb = RGBColor(*self._docx_report_settings.heading_color_rgb)
            run.font.bold = True
            run.font.underline = True

        document.add_paragraph()

        document_title = report_input.document_content.title or report_input.filename
        title_paragraph = document.add_paragraph()
        title_paragraph.add_run(document_title).bold = True
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_paragraph.runs[0].font.size = Pt(self._docx_report_settings.base_font_size_pt)

    def _add_header_logo(self, document: DocxDocument) -> None:
        if self._logo_path is None:
            return
        section = document.sections[0]
        header = section.header
        header.paragraphs[0].clear()

        table = header.add_table(
            rows=1,
            cols=2,
            width=Inches(self._docx_report_settings.header_table_width_inches),
        )

        left_cell = table.rows[0].cells[0]
        left_cell.width = Inches(self._docx_report_settings.header_left_cell_width_inches)
        left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        left_paragraph = left_cell.paragraphs[0]
        left_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        app_label_run = left_paragraph.add_run(
            f"Generado por {self._docx_report_settings.app_name} v{self._docx_report_settings.app_version}\n"
        )
        app_label_run.italic = True
        app_label_run.font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
        app_label_run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

        date_run = left_paragraph.add_run(self._clock().strftime("%d de %B de %Y"))
        date_run.font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
        date_run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

        right_cell = table.rows[0].cells[1]
        right_cell.width = Inches(self._docx_report_settings.header_right_cell_width_inches)
        right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right_paragraph = right_cell.paragraphs[0]
        right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        self._add_logo_picture(paragraph=right_paragraph, logo_path=self._logo_path)

        table_element = table._element
        table_properties = table_element.tblPr
        if table_properties is None:
            table_properties = OxmlElement("w:tblPr")
            table_element.insert(0, table_properties)

        table_borders = OxmlElement("w:tblBorders")
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "none")
            border.set(qn("w:sz"), "0")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "auto")
            table_borders.append(border)
        table_properties.append(table_borders)

    def _add_logo_picture(self, paragraph: Paragraph, logo_path: str) -> None:
        if not (isfile(logo_path) and access(logo_path, R_OK)):
            logger.warning("Header logo is not readable, skipping picture: %s", logo_path)
            return
        logo_run = paragraph.add_run()
        logo_run.add_picture(logo_path, width=Inches(self._docx_report_settings.logo_width_inches))

    def _add_page_numbers(self, document: DocxDocument) -> None:
        section = document.sections[0]
        footer = section.footer
        footer.paragraphs[0].clear()

        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        page_number_run = paragraph.add_run()

        page_field_begin = OxmlElement("w:fldChar")
        page_field_begin.set(qn("w:fldCharType"), "begin")
        page_number_run._element.append(page_field_begin)

        page_instruction = OxmlElement("w:instrText")
        page_instruction.set(qn("xml:space"), "preserve")
        page_instruction.text = "PAGE"
        page_number_run._element.append(page_instruction)

        page_field_end = OxmlElement("w:fldChar")
        page_field_end.set(qn("w:fldCharType"), "end")
        page_number_run._element.append(page_field_end)

        separator_run = paragraph.add_run(" de ")
        separator_run.font.size = Pt(self._docx_report_settings.page_number_font_size_pt)

        total_pages_run = paragraph.add_run()

        total_pages_field_begin = OxmlElement("w:fldChar")
        total_pages_field_begin.set(qn("w:fldCharType"), "begin")
        total_pages_run._element.append(total_pages_field_begin)

        total_pages_instruction = OxmlElement("w:instrText")
        total_pages_instruction.set(qn("xml:space"), "preserve")
        total_pages_instruction.text = "NUMPAGES"
        total_pages_run._element.append(total_pages_instruction)

        total_pages_field_end = OxmlElement("w:fldChar")
        total_pages_field_end.set(qn("w:fldCharType"), "end")
        total_pages_run._element.append(total_pages_field_end)

        for text_run in paragraph.runs:
            text_run.font.size = Pt(self._docx_report_settings.page_number_font_size_pt)
            text_run.font.color.rgb = RGBColor(*self._docx_report_settings.neutral_color_rgb)

    def _add_executive_summary(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(document=document, text="RESUMEN EJECUTIVO", level=1)

        can_publish = report_input.is_publishable
        reason = report_input.publishability_reason

        decision_paragraph = document.add_paragraph()
        decision_run = decision_paragraph.add_run(
            "✅ APTO PARA PUBLICACIÓN" if can_publish else "⚠️ REQUIERE REVISIÓN"
        )
        decision_run.bold = True
        decision_run.font.size = Pt(self._docx_report_settings.decision_font_size_pt)
        decision_run.font.color.rgb = (
            RGBColor(*self._docx_report_settings.publishable_color_rgb)
            if can_publish
            else RGBColor(*self._docx_report_settings.reject_color_rgb)
        )

        document.add_paragraph(reason)
        document.add_paragraph()

        table = document.add_table(rows=6, cols=2)
        table.style = self._docx_report_settings.table_style

        header_cells = table.rows[0].cells
        header_cells[0].text = "Métrica"
        header_cells[1].text = "Valor"
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        apa_count = len(report_input.apa_validation.violations)

        table.rows[1].cells[0].text = "Calidad General"
        table.rows[1].cells[1].text = f"{report_input.quality.overall_score:.1f}/10"

        table.rows[2].cells[0].text = "Gramática y Ortografía"
        table.rows[2].cells[1].text = f"{report_input.grammar.score:.1f}/10"

        table.rows[3].cells[0].text = "Estructura"
        table.rows[3].cells[1].text = (
            "✓ Válida" if report_input.structure.is_valid else "✗ Incompleta"
        )

        table.rows[4].cells[0].text = "Errores APA 7"
        table.rows[4].cells[1].text = f"{apa_count} detectados" if apa_count > 0 else "Sin errores"

        table.rows[5].cells[0].text = "Tasa de Coincidencia"
        table.rows[5].cells[1].text = self._citation_match_rate_formatter.format(
            citations=report_input.citations
        )

        document.add_paragraph()

    def _add_document_info(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="📄 INFORMACIÓN DEL DOCUMENTO", level=1
        )

        document_content = report_input.document_content
        estimated_pages = document_content.word_count // self._docx_report_settings.words_per_page

        self._labeled_paragraph_writer.write(
            document=document,
            label="Título: ",
            value=document_content.title or report_input.filename,
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Autor(es): ",
            value=document_content.authors or "No identificado",
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Total de palabras: ",
            value=f"{document_content.word_count:,}",
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Total de caracteres: ",
            value=f"{document_content.char_count:,}",
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Páginas estimadas: ",
            value=str(estimated_pages),
        )

    def _add_classification(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="🏷️ CLASIFICACIÓN DEL ARTÍCULO", level=1
        )

        classification = report_input.classification

        self._labeled_paragraph_writer.write(
            document=document,
            label="Categoría: ",
            value=classification.article_type.value.upper(),
        )

        confidence_text = (
            f"{classification.confidence:.1%}" if classification.confidence is not None else "—"
        )
        self._labeled_paragraph_writer.write(
            document=document,
            label="Confianza: ",
            value=confidence_text,
        )

        if classification.reasoning:
            self._labeled_paragraph_writer.write(
                document=document,
                label="Razonamiento: ",
                value=classification.reasoning,
            )

    def _add_quality_analysis(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="⭐ ANÁLISIS DE CALIDAD SEMÁNTICA", level=1
        )

        quality = report_input.quality

        paragraph = document.add_paragraph()
        paragraph.add_run("Puntuación General: ").bold = True
        score_run = paragraph.add_run(f"{quality.overall_score:.1f}/10")
        score_run.bold = True
        score_run.font.size = Pt(self._docx_report_settings.score_font_size_pt)
        score_run.font.color.rgb = RGBColor(*self._color_for_score(score=quality.overall_score))

        if quality.dimension_scores:
            for dimension_name, dimension_data in quality.dimension_scores.items():
                document.add_heading(QualityDimension.label_for(dimension_name), level=3)

                self._labeled_paragraph_writer.write(
                    document=document,
                    label="Puntuación: ",
                    value=f"{dimension_data['score']:.1f}/10",
                )

                feedback_blocks = dimension_data.get("feedback_blocks")
                if feedback_blocks:
                    self._feedback_block_writer.write_blocks(
                        document=document, blocks=feedback_blocks
                    )
                elif dimension_data.get("feedback"):
                    self._markdown_run_writer.write_paragraph(
                        document=document,
                        text=dimension_data["feedback"],
                    )

    def _add_editorial_suitability(
        self, document: DocxDocument, report_input: ReportInputDTO
    ) -> None:
        suitability = report_input.quality.editorial_suitability
        if suitability is None:
            return

        self._section_heading_writer.write(
            document=document, text="🎯 PERTINENCIA EDITORIAL", level=1
        )

        document.add_heading("Contribución", level=3)
        self._labeled_paragraph_writer.write(
            document=document,
            label="Veredicto: ",
            value=suitability.contribution_verdict,
        )

        if suitability.contribution_phrase:
            self._labeled_paragraph_writer.write(
                document=document,
                label="Aporte identificado: ",
                value=suitability.contribution_phrase,
            )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Observación: ",
            value=suitability.contribution_observation,
        )

        document.add_heading("Alineación con líneas de investigación", level=3)
        self._labeled_paragraph_writer.write(
            document=document,
            label="Veredicto: ",
            value=suitability.alignment_verdict,
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Líneas relacionadas: ",
            value=suitability.alignment_lines,
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Justificación: ",
            value=suitability.alignment_justification,
        )

    def _add_grammar_analysis(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="📝 GRAMÁTICA Y ORTOGRAFÍA", level=1
        )

        grammar = report_input.grammar

        self._labeled_paragraph_writer.write(
            document=document,
            label="Puntuación: ",
            value=f"{grammar.score:.1f}/10",
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Estado: ",
            value=grammar.feedback,
        )

        if grammar.errors:
            document.add_paragraph()
            document.add_paragraph("Errores Detectados:").bold = True

            max_errors = self._docx_report_settings.max_errors_displayed
            context_limit = self._docx_report_settings.context_truncation_limit
            for grammar_error in grammar.errors[:max_errors]:
                document.add_paragraph(grammar_error.message, style="List Number")

                context_text = (
                    grammar_error.context
                    if len(grammar_error.context) < context_limit
                    else grammar_error.context[:context_limit] + "..."
                )
                document.add_paragraph(f'   Contexto: "{context_text}"', style="List Bullet 2")

                if grammar_error.replacements:
                    max_replacements = self._docx_report_settings.max_replacements
                    document.add_paragraph(
                        f"   Sugerencia: {', '.join(grammar_error.replacements[:max_replacements])}",
                        style="List Bullet 2",
                    )

    def _add_apa_validation(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="📖 VALIDACIÓN APA 7 (ESPAÑOL)", level=1
        )

        violations = report_input.apa_validation.violations

        if not violations:
            no_citations = report_input.citations.total_citations == 0
            message = (
                "ℹ️ Sin citaciones en texto detectadas — no se pudo validar formato APA 7"
                if no_citations
                else "✅ Sin errores de formato APA 7 detectados"
            )
            color = (
                self._docx_report_settings.neutral_color_rgb
                if no_citations
                else self._docx_report_settings.publishable_color_rgb
            )
            document.add_paragraph().add_run(message).font.color.rgb = RGBColor(*color)
            document.add_paragraph()
            return

        violations_by_type = defaultdict(list)
        for violation in violations:
            violations_by_type[violation.error_type].append(violation)

        max_errors = self._docx_report_settings.max_errors_displayed
        for _, errors in violations_by_type.items():
            document.add_paragraph(f"CITACIÓN INCORRECTA ({len(errors)}):", style="Heading 3")

            for index, apa_violation in enumerate(errors[:max_errors], 1):
                document.add_paragraph(f"{index}. Citación: {apa_violation.citation_text}")
                document.add_paragraph(
                    f'   Ubicación: "{apa_violation.location}"', style="List Bullet 2"
                )
                document.add_paragraph(
                    f"   Problema: {apa_violation.explanation}", style="List Bullet 2"
                )
                if apa_violation.correction:
                    document.add_paragraph(
                        f"   Corrección: {apa_violation.correction}", style="List Bullet 2"
                    )

    def _add_structure_validation(
        self, document: DocxDocument, report_input: ReportInputDTO
    ) -> None:
        self._section_heading_writer.write(
            document=document, text="📋 VALIDACIÓN DE ESTRUCTURA (EUMIC)", level=1
        )

        structure = report_input.structure

        if structure.is_valid:
            document.add_paragraph().add_run(
                "✓ Estructura válida según normas EUMIC"
            ).font.color.rgb = RGBColor(*self._docx_report_settings.publishable_color_rgb)
            return

        document.add_paragraph().add_run("✗ Estructura incompleta").font.color.rgb = RGBColor(
            *self._docx_report_settings.reject_color_rgb
        )

        if not structure.missing_sections:
            return

        document.add_paragraph()
        document.add_paragraph("Secciones faltantes:").bold = True
        for section in structure.missing_sections:
            document.add_paragraph(f"• {section}", style="List Bullet")

    def _add_citations_analysis(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(
            document=document, text="📚 ANÁLISIS DE CITAS Y REFERENCIAS", level=1
        )

        citations = report_input.citations

        table = document.add_table(rows=4, cols=2)
        table.style = self._docx_report_settings.table_style

        table.rows[0].cells[0].text = "Total de citas en texto"
        table.rows[0].cells[1].text = str(citations.total_citations)

        table.rows[1].cells[0].text = "Total de referencias bibliográficas"
        table.rows[1].cells[1].text = str(citations.total_references)

        table.rows[2].cells[0].text = "Citas con referencia"
        table.rows[2].cells[1].text = str(citations.matched_count)

        table.rows[3].cells[0].text = "Tasa de coincidencia"
        table.rows[3].cells[1].text = self._citation_match_rate_formatter.format(
            citations=citations
        )

    def _add_recommendations(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        self._section_heading_writer.write(document=document, text="💡 RECOMENDACIONES", level=1)

        verdict = report_input.verdict
        verdict_colors = {
            PublicationVerdict.CRITICAL: self._docx_report_settings.reject_color_rgb,
            PublicationVerdict.WARNING: self._docx_report_settings.reject_color_rgb,
            PublicationVerdict.APPROVED: self._docx_report_settings.publishable_color_rgb,
        }
        paragraph = document.add_paragraph()
        paragraph.add_run(verdict.message).bold = True
        paragraph.runs[0].font.size = Pt(self._docx_report_settings.recommendation_font_size_pt)
        paragraph.runs[0].font.color.rgb = RGBColor(*verdict_colors[verdict.verdict])

        if not report_input.recommendations:
            return

        priority_icons = {
            RecommendationPriority.HIGH: "🔴",
            RecommendationPriority.MEDIUM: "🟡",
            RecommendationPriority.LOW: "🟢",
        }
        document.add_paragraph("Recomendaciones específicas:").bold = True
        for recommendation in report_input.recommendations:
            icon = priority_icons.get(recommendation.priority, "⚪")
            document.add_paragraph(f"{icon} {recommendation.message}", style="List Bullet")

    def _add_footer(self, document: DocxDocument) -> None:
        document.add_paragraph()
        document.add_paragraph("─" * 80)

        footer_paragraph = document.add_paragraph()
        footer_paragraph.add_run(
            f"Generado por {self._docx_report_settings.app_name} v{self._docx_report_settings.app_version} | "
        ).italic = True
        footer_paragraph.add_run(self._clock().strftime("%d/%m/%Y %H:%M")).italic = True
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_paragraph.runs[0].font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
        footer_paragraph.runs[0].font.color.rgb = RGBColor(
            *self._docx_report_settings.neutral_color_rgb
        )
        footer_paragraph.runs[1].font.size = Pt(self._docx_report_settings.metadata_font_size_pt)
        footer_paragraph.runs[1].font.color.rgb = RGBColor(
            *self._docx_report_settings.neutral_color_rgb
        )

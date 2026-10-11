from collections import defaultdict

from docx.document import Document as DocxDocument
from docx.shared import RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)


class DocxApaValidationSection(DocxReportSection):
    """Render the APA 7 citation format violations grouped by error type."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add a colored no-error notice or each violation group with its details."""
        self._section_heading_writer.write(
            document=document, text="📖 VALIDACIÓN APA 7 (ESPAÑOL)", level=1
        )

        violations = report_input.apa_validation.violations

        if not violations:
            no_citations = report_input.citations.total_citations == 0
            message = (
                "ℹ️ Sin citaciones en texto detectadas — no se pudo validar formato APA 7"
                if no_citations
                else "✅ Sin errores de formato APA 7 detectados"
            )
            color = (
                self._docx_report_settings.neutral_color_rgb
                if no_citations
                else self._docx_report_settings.publishable_color_rgb
            )
            document.add_paragraph().add_run(message).font.color.rgb = RGBColor(*color)
            document.add_paragraph()
            return

        violations_by_type = defaultdict(list)
        for violation in violations:
            violations_by_type[violation.error_type].append(violation)

        max_errors = self._docx_report_settings.max_errors_displayed
        for errors in violations_by_type.values():
            document.add_paragraph(f"CITACIÓN INCORRECTA ({len(errors)}):", style="Heading 3")

            for index, apa_violation in enumerate(errors[:max_errors], 1):
                document.add_paragraph(f"{index}. Citación: {apa_violation.citation_text}")
                document.add_paragraph(
                    f'   Ubicación: "{apa_violation.location}"', style="List Bullet 2"
                )
                document.add_paragraph(
                    f"   Problema: {apa_violation.explanation}", style="List Bullet 2"
                )
                if apa_violation.correction:
                    document.add_paragraph(
                        f"   Corrección: {apa_violation.correction}", style="List Bullet 2"
                    )

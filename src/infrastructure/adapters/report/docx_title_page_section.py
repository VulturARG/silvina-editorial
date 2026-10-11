from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings

REPORT_TITLE_HEADING_LEVEL = 0


class DocxTitlePageSection(DocxReportSection):
    """Render the report title and the analyzed document title."""

    def __init__(self, docx_report_settings: DocxReportSettings) -> None:
        self._docx_report_settings = docx_report_settings

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the centered report title followed by the analyzed document title."""
        title = document.add_heading("INFORME DE ANÁLISIS EDITORIAL", REPORT_TITLE_HEADING_LEVEL)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.size = Pt(self._docx_report_settings.title_font_size_pt)
            run.font.color.rgb = RGBColor(*self._docx_report_settings.heading_color_rgb)
            run.font.bold = True
            run.font.underline = True

        document.add_paragraph()

        document_title = report_input.document_content.title or report_input.filename
        title_paragraph = document.add_paragraph()
        title_paragraph.add_run(document_title).bold = True
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_paragraph.runs[0].font.size = Pt(self._docx_report_settings.base_font_size_pt)

from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.citation_match_rate_formatter import (
    CitationMatchRateFormatter,
)
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)

METRICS_TABLE_ROW_COUNT = 6
METRICS_TABLE_COLUMN_COUNT = 2


class DocxExecutiveSummarySection(DocxReportSection):
    """Render the publication decision, its reason and the key metrics table."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
        citation_match_rate_formatter: CitationMatchRateFormatter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer
        self._citation_match_rate_formatter = citation_match_rate_formatter

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the colored decision, the publishability reason and the metrics table."""
        self._section_heading_writer.write(document=document, text="RESUMEN EJECUTIVO", level=1)

        publish_threshold = self._docx_report_settings.publish_threshold
        can_publish = report_input.is_publishable(publish_threshold)
        reason = report_input.publishability_reason(publish_threshold)

        decision_paragraph = document.add_paragraph()
        decision_run = decision_paragraph.add_run(
            "✅ APTO PARA PUBLICACIÓN" if can_publish else "⚠️ REQUIERE REVISIÓN"
        )
        decision_run.bold = True
        decision_run.font.size = Pt(self._docx_report_settings.decision_font_size_pt)
        decision_run.font.color.rgb = (
            RGBColor(*self._docx_report_settings.publishable_color_rgb)
            if can_publish
            else RGBColor(*self._docx_report_settings.reject_color_rgb)
        )

        document.add_paragraph(reason)
        document.add_paragraph()

        table = document.add_table(rows=METRICS_TABLE_ROW_COUNT, cols=METRICS_TABLE_COLUMN_COUNT)
        table.style = self._docx_report_settings.table_style

        header_cells = table.rows[0].cells
        header_cells[0].text = "Métrica"
        header_cells[1].text = "Valor"
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        apa_count = len(report_input.apa_validation.violations)

        table.rows[1].cells[0].text = "Calidad General"
        table.rows[1].cells[1].text = f"{report_input.quality.overall_score:.1f}/10"

        table.rows[2].cells[0].text = "Gramática y Ortografía"
        table.rows[2].cells[1].text = f"{report_input.grammar.score:.1f}/10"

        table.rows[3].cells[0].text = "Estructura"
        table.rows[3].cells[1].text = (
            "✓ Válida" if report_input.structure.is_valid else "✗ Incompleta"
        )

        table.rows[4].cells[0].text = "Errores APA 7"
        table.rows[4].cells[1].text = f"{apa_count} detectados" if apa_count > 0 else "Sin errores"

        table.rows[5].cells[0].text = "Tasa de Coincidencia"
        table.rows[5].cells[1].text = self._citation_match_rate_formatter.format(
            citations=report_input.citations
        )

        document.add_paragraph()

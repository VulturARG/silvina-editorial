from collections.abc import Iterator
from re import Pattern
from re import compile as regex_compile

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

_MARKDOWN_BOLD_PATTERN = regex_compile(r"\*\*(.+?)\*\*(?!\*)")
_MARKDOWN_ITALIC_PATTERN = regex_compile(r"(?<!\*)\*(?!\s|\*)([^\r\n*]+?)(?<!\s)\*(?!\*)")


class DocxMarkdownRunWriter:
    """Render Markdown bold and italic segments into Word document paragraphs."""

    def write_paragraph(
        self,
        document: DocxDocument,
        text: str,
        paragraph: Paragraph | None = None,
        prefix: str = "",
    ) -> Paragraph:
        """Add or populate a paragraph rendering bold and italic Markdown segments."""
        target_paragraph = paragraph if paragraph is not None else document.add_paragraph()
        self._append_prefix(target_paragraph, prefix)
        for segment_text, is_bold in self._iter_pattern_segments(text, _MARKDOWN_BOLD_PATTERN):
            self._render_italic_segments(target_paragraph, segment_text, is_bold=is_bold)
        return target_paragraph

    def _render_italic_segments(
        self,
        paragraph: Paragraph,
        text: str,
        is_bold: bool,
    ) -> None:
        for segment_text, is_italic in self._iter_pattern_segments(text, _MARKDOWN_ITALIC_PATTERN):
            self._add_styled_run(
                paragraph=paragraph,
                text=segment_text,
                is_bold=is_bold,
                is_italic=is_italic,
            )

    def _append_prefix(self, paragraph: Paragraph, prefix: str) -> None:
        if not prefix:
            return
        paragraph.add_run(prefix)

    def _add_styled_run(
        self,
        paragraph: Paragraph,
        text: str,
        is_bold: bool = False,
        is_italic: bool = False,
    ) -> None:
        run = paragraph.add_run(text)
        if is_bold:
            run.bold = True
        if is_italic:
            run.italic = True

    def _iter_pattern_segments(
        self,
        text: str,
        pattern: Pattern[str],
    ) -> Iterator[tuple[str, bool]]:
        position = 0
        for match in pattern.finditer(text):
            if match.start() > position:
                yield (text[position : match.start()], False)
            yield (match.group(1), True)
            position = match.end()
        if position < len(text):
            yield (text[position:], False)

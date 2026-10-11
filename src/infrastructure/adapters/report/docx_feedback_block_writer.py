from collections.abc import Callable, Sequence

from docx.document import Document as DocxDocument
from docx.shared import Inches

from src.domain.enums.feedback_block_kind import FeedbackBlockKind
from src.infrastructure.adapters.report.docx_feedback_block_payload import (
    DocxFeedbackBlockPayload,
)
from src.infrastructure.adapters.report.docx_feedback_block_view import (
    DocxFeedbackBlockView,
)
from src.infrastructure.adapters.report.docx_markdown_run_writer import (
    DocxMarkdownRunWriter,
)
from src.infrastructure.adapters.report.docx_report_settings import (
    DocxReportSettings,
)


class DocxFeedbackBlockWriter:
    """Render structured feedback blocks into a Word document."""

    def __init__(
        self,
        markdown_run_writer: DocxMarkdownRunWriter,
        docx_report_settings: DocxReportSettings,
    ) -> None:
        self._markdown_run_writer = markdown_run_writer
        self._docx_report_settings = docx_report_settings
        self._writers: dict[
            FeedbackBlockKind,
            Callable[[DocxDocument, DocxFeedbackBlockView], None],
        ] = {
            FeedbackBlockKind.TITLE: self._write_title,
            FeedbackBlockKind.ITEM: self._write_item,
            FeedbackBlockKind.TEXT: self._write_text,
        }

    def write_blocks(
        self,
        document: DocxDocument,
        blocks: Sequence[DocxFeedbackBlockPayload],
    ) -> None:
        """Add structured feedback blocks to the document."""
        for block in blocks:
            view = DocxFeedbackBlockView.from_payload(block)
            self._writers[view.kind](document, view)

    def _write_title(
        self,
        document: DocxDocument,
        view: DocxFeedbackBlockView,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run(view.text).bold = True

    def _write_item(
        self,
        document: DocxDocument,
        view: DocxFeedbackBlockView,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(self._calculate_indent_inches(view.level))
        self._markdown_run_writer.write_paragraph(
            document=document,
            text=view.text,
            paragraph=paragraph,
            prefix=self._format_marker_prefix(view.marker),
        )

    def _write_text(
        self,
        document: DocxDocument,
        view: DocxFeedbackBlockView,
    ) -> None:
        self._markdown_run_writer.write_paragraph(
            document=document,
            text=view.text,
        )

    def _calculate_indent_inches(self, level: int) -> float:
        indents = self._docx_report_settings.feedback_item_indent_inches
        clamped_index = min(max(level, 0), len(indents) - 1)
        return indents[clamped_index]

    def _format_marker_prefix(self, marker: str) -> str:
        if not marker:
            return ""
        return marker if marker.endswith(" ") else f"{marker} "

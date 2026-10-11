from docx.document import Document as DocxDocument

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.infrastructure.adapters.report.docx_labeled_paragraph_writer import (
    DocxLabeledParagraphWriter,
)
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)

TRUNCATION_SUFFIX = "..."


class DocxGrammarAnalysisSection(DocxReportSection):
    """Render the grammar score, its status and the detected grammar errors."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
        labeled_paragraph_writer: DocxLabeledParagraphWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer
        self._labeled_paragraph_writer = labeled_paragraph_writer

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the score and status, then list errors with context and suggestions."""
        self._section_heading_writer.write(
            document=document, text="📝 GRAMÁTICA Y ORTOGRAFÍA", level=1
        )

        grammar = report_input.grammar

        self._labeled_paragraph_writer.write(
            document=document,
            label="Puntuación: ",
            value=f"{grammar.score:.1f}/10",
        )

        self._labeled_paragraph_writer.write(
            document=document,
            label="Estado: ",
            value=grammar.feedback,
        )

        if grammar.errors:
            document.add_paragraph()
            document.add_paragraph().add_run("Errores Detectados:").bold = True

            max_errors = self._docx_report_settings.max_errors_displayed
            context_limit = self._docx_report_settings.context_truncation_limit
            for grammar_error in grammar.errors[:max_errors]:
                document.add_paragraph(grammar_error.message, style="List Number")

                context_text = (
                    grammar_error.context
                    if len(grammar_error.context) < context_limit
                    else grammar_error.context[:context_limit] + TRUNCATION_SUFFIX
                )
                document.add_paragraph(f'   Contexto: "{context_text}"', style="List Bullet 2")

                if grammar_error.replacements:
                    max_replacements = self._docx_report_settings.max_replacements
                    document.add_paragraph(
                        f"   Sugerencia: {', '.join(grammar_error.replacements[:max_replacements])}",
                        style="List Bullet 2",
                    )

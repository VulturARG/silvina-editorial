from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor

from src.domain.dtos.report_input_dto import ReportInputDTO
from src.domain.enums.publication_verdict import PublicationVerdict
from src.domain.enums.recommendation_priority import RecommendationPriority
from src.infrastructure.adapters.report.docx_report_section import DocxReportSection
from src.infrastructure.adapters.report.docx_report_settings import DocxReportSettings
from src.infrastructure.adapters.report.docx_section_heading_writer import (
    DocxSectionHeadingWriter,
)

PRIORITY_ICONS = {
    RecommendationPriority.HIGH: "🔴",
    RecommendationPriority.MEDIUM: "🟡",
    RecommendationPriority.LOW: "🟢",
}
UNKNOWN_PRIORITY_ICON = "⚪"


class DocxRecommendationsSection(DocxReportSection):
    """Render the publication verdict message and the prioritized recommendations."""

    def __init__(
        self,
        docx_report_settings: DocxReportSettings,
        section_heading_writer: DocxSectionHeadingWriter,
    ) -> None:
        self._docx_report_settings = docx_report_settings
        self._section_heading_writer = section_heading_writer
        self._verdict_colors = {
            PublicationVerdict.CRITICAL: docx_report_settings.reject_color_rgb,
            PublicationVerdict.WARNING: docx_report_settings.reject_color_rgb,
            PublicationVerdict.APPROVED: docx_report_settings.publishable_color_rgb,
        }

    def render(self, document: DocxDocument, report_input: ReportInputDTO) -> None:
        """Add the verdict colored by severity and list recommendations by priority icon."""
        self._section_heading_writer.write(document=document, text="💡 RECOMENDACIONES", level=1)

        verdict = report_input.verdict
        paragraph = document.add_paragraph()
        paragraph.add_run(verdict.message).bold = True
        paragraph.runs[0].font.size = Pt(self._docx_report_settings.recommendation_font_size_pt)
        paragraph.runs[0].font.color.rgb = RGBColor(*self._verdict_colors[verdict.verdict])

        if not report_input.recommendations:
            return

        document.add_paragraph().add_run("Recomendaciones específicas:").bold = True
        for recommendation in report_input.recommendations:
            icon = PRIORITY_ICONS.get(recommendation.priority, UNKNOWN_PRIORITY_ICON)
            document.add_paragraph(f"{icon} {recommendation.message}", style="List Bullet")
